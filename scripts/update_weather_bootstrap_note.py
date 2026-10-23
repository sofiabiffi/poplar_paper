from pathlib import Path
from datetime import datetime
from copy import deepcopy
import shutil
from docx import Document
from docx.text.paragraph import Paragraph

root = Path(r'C:/Users/sofia/Documents/poplar_paper')
path = root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
history = path.parent/'_history'
history.mkdir(exist_ok=True)
shutil.copy2(path, history/f'02_weather_and_PAIC_before_bootstrap_{datetime.now():%Y%m%d_%H%M%S}.docx')
d = Document(path)
start = next(p for p in d.paragraphs if p.text in ('What the dots and whiskers mean', 'Bootstrap uncertainty'))
end = next(p for p in d.paragraphs if p.text == 'Calculated weather associations by hybrid')
body_template = next(p for p in d.paragraphs if p.text.startswith(('Dots are estimated', 'Bootstrapping asks')))
heading_template = deepcopy(start._p)
body_properties = deepcopy(body_template._p.pPr)
node = start._p
while node is not end._p:
    following = node.getnext()
    node.getparent().remove(node)
    node = following

def insert(text, heading=False):
    p = end.insert_paragraph_before(text)
    if heading:
        p.style = 'Heading 1'
    else:
        if p._p.pPr is not None:
            p._p.remove(p._p.pPr)
        if body_properties is not None:
            p._p.insert(0, deepcopy(body_properties))
    return p

p = insert('Bootstrap uncertainty', True)
p.paragraph_format.page_break_before = True
insert('A simple explanation') .style = 'Heading 2'
insert('Bootstrapping asks how much the estimated weather association changes when the observed record is given different emphasis. We repeatedly draw observations from that record, allowing some to appear more than once and others to be left out, then fit the model again. We do not invent new growth measurements or new weather values.')
insert('Here, we sample short blocks of three successive measurement occasions, rather than isolated years, because nearby observations may be related. Each selected occasion carries its PAIC, age and assigned weather together, and the same occasions are selected for all four hybrids. This retains their shared weather history and some temporal dependence. A block contains three observed visits, which can span more than three calendar years when measurements are missing. Blocks can wrap from the end of the record to the beginning; this is the circular part of the procedure.')
insert('We repeat this 1,000 times. Each time, we refit all four hybrid models, including their quadratic age terms, and recalculate the absolute PAIC differences for 1 °C higher summer temperature and 100 mm more summer rainfall. The other weather variable is held at its study mean, and differences are averaged over the same 26 observed ages. The central 95% of the 1,000 estimates forms each whisker. A wide whisker means the average association is less precisely estimated; it does not show how the biological response changes from year to year.')
insert('Suggested wording for the paper', True)
insert('Uncertainty in absolute weather associations was estimated using 1,000 paired circular block-bootstrap resamples. Blocks of three successive measurement occasions were sampled with replacement, retaining PAIC, age and assigned weather together and using the same sampled occasions for all four hybrids. Blocks contained observed visits rather than necessarily consecutive calendar years. For each resample, the hybrid-specific models, including quadratic age, were refitted and absolute mean PAIC contrasts for a 1 °C increase in summer temperature and a 100 mm increase in summer rainfall were recalculated. Contrasts were averaged over the same 26 observed ages, holding the other weather predictor at its study mean. The 2.5th and 97.5th percentiles of the resampled estimates provided pointwise 95% uncertainty intervals. Differences between hybrids were calculated within each resample.')
insert('Interpretation and limits', True)
insert('Pairing allows direct comparisons between hybrids under the same resampled weather record. The blocks preserve some dependence, but three visits are not established as sufficient to capture all temporal dependence. These exploratory intervals are conditional on the selected summer model and do not include model-selection, allometric or weather measurement uncertainty. They are pointwise intervals, without adjustment for multiple comparisons. Extreme resampled estimates were retained; their diagnostic caveat is described in the Discussion.')
d.save(path)
print(path)
