from pathlib import Path
from datetime import datetime
import shutil,csv
from docx import Document
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/04_soil_carbon_stocks.docx'
backup=path.parent/'_history'/('04_soil_carbon_stocks_before_provenance_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx')
backup.parent.mkdir(parents=True, exist_ok=True); shutil.copy2(path,backup)
d=Document(path)
# Preserve existing prose and figures; add a conspicuous qualification near the front.
p=d.paragraphs[2].insert_paragraph_before('Provenance update 22 October 2026: all 188 measured records have been traced to AgricARBON sample IDs. Tree row profile 25 matches reference-field core 016. Earlier numerical summaries and tests below retain the original workbook labels and require revision before scenario use. See the provenance addendum at the end.')
p.runs[0].bold=True
d.add_page_break()
d.add_heading('Soil sample provenance and former cropping treatments',level=1)
d.add_paragraph('Updated 22 October 2026. The missing sample identities can be recovered. The three agroforestry sampling bands are consistent with the colleague\'s description of rows 6, 10 and 14, but their geographic numbering remains provisional until an original site anchor is confirmed.')
d.add_heading('Methods',level=2)
d.add_paragraph('Match rawdataSofia records to CSV - Samples in AGRICARBON UOL001-A01 Results v3.3 PUBLISH 2024-03-14 CSV.xlsx using fine dry bulk density and SOC percentage, with a tolerance of 0.005001 in each value to allow two-decimal rounding. Sample suffix .1 is aligned with 0-15 cm and .2 with 15-30 cm. Intersect candidate core IDs across both available layers; retain the original sampling-location codes and potential_samplingID. Inspect page 4 of UOL001 Field Plan.pdf and the original experimental-design screenshot supplied with the request. Geographic bands are inferred from coordinates; treatment row numbers use the colleague\'s account and the screenshot\'s west-to-east ordering provisionally.')
d.add_heading('Calculated results',level=2)
d.add_paragraph('All 188 records match both BD and SOC within rounding tolerance. Each of the 96 profile fingerprints resolves to one distinct AgricARBON core. There are 48 cores in agroforestry Field 2: 24 labelled Alley and 24 labelled Tree row. The remaining 48 cores are in reference Field 1: 47 labelled Arable reference and one incorrectly labelled Tree row. This is a provenance finding, not a recalculation of SOC stocks.')
t=d.add_table(rows=1,cols=4);t.style='Table Grid'
for c,txt in zip(t.rows[0].cells,['Spatial band','Alley workbook IDs','Provisional row and block','Former treatment']):c.text=txt
for vals in [('Western','1-8','Row 6 / Block 1','Continuous fallow'),('Middle','9-16','Row 10 / Block 2','Continuous cropping'),('Eastern','17-24','Row 14 / Block 3','Alternate cropping')]:
    for c,txt in zip(t.add_row().cells,vals):c.text=txt
d.add_paragraph('These are potential_samplingID values within 2_alley, not company core numbers. Each group contains eight profiles with both measured layers. Permanent core IDs and coordinates are retained in results/checks/soil_provenance/soil_profile_crosswalk.csv. The row and treatment columns in that file are explicitly provisional.')
d.add_paragraph('Tree row profile 25 (1_tree_25) matches UOL001-A01-001.0-A-016 at latitude 53.873887, longitude -1.333282, in Field 1. Both layers match. The corresponding profile is absent from the 47 profiles labelled Arable reference. A separate case, Tree row profile 12, matches Field 2 core 032 and lies in the middle band; core-number ranges alone must not define sampling location.')
d.add_heading('Discussion hypotheses',level=2)
d.add_paragraph('Within this soil subset each former cropping treatment is represented in only one block and one sampled tree-row band. Multiple cores within a band are spatial subsamples, not independent block replication of management. Treatment differences cannot be separated from block and band differences using these data alone. Alternate cropping denotes the historical rotation in which alley positions changed annually; the sampled side cannot establish a permanent Continuous fallow or Continuous cropping history. Exact hybrid segments and the crop status of a side in a particular year remain unassigned.')
d.add_heading('Open decisions',level=2)
d.add_paragraph('Confirm that the western, middle and eastern sampled bands are rows 6, 10 and 14. Page 4 links to a separate agroforestry_sampling_design.pdf on Agricarbon SharePoint; that file is not in the supplied sources and may resolve the anchor. Confirm reference-field management: the plan calls Field 1 open grassland, whereas the scenario workbook calls it arable. The plan is labelled Spring 2024 and date extracted is 24-25 January 2024 for the matched records; determine whether the stated 2023 date refers to a different operation. Do not treat date extracted as a separately verified laboratory or field event.')
d.add_paragraph('Before SOC scenario calculations, resolve the reference-field label and reassign Tree row profile 25 in the analysis inputs, then rerun fixed-depth stocks, ESM reference masses, summaries and tests. Existing outputs have not been silently changed. A cross-sectional stock contrast does not by itself identify sequestration through time. Reproduce the identity audit with scripts/trace_soil_provenance.py; the value audit retains all original and matched values.')
d.save(path)
print(path)
