from pathlib import Path
from shutil import copy2
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
path=root/"manuscript/analysis_notes/03_cropping_treatment_and_legacy.docx"
backup=path.parent/"_qa"/"03_cropping_treatment_and_legacy"/"before_2026_10_22_update.docx"
backup.parent.mkdir(parents=True,exist_ok=True)
if not backup.exists(): copy2(path,backup)
d=Document(path)
def insert(anchor,text,style="Normal"):
 p=anchor.insert_paragraph_before(text,style)
 return p
anchor=next(p for p in d.paragraphs if p.text=="Main results")
insert(anchor,"Model specification and block adjustment","Heading 1")
insert(anchor,"The following pseudo-R specification mirrors analysis/03_treatment_active.R. Fit it separately to the active and post-cropping subsets, and repeat with the boundary interval included in active cropping for sensitivity. PAIC is the equally weighted mean across available measured trees within each row, hybrid and interval. Block, treatment, hybrid and interval are factors.")
code='''# IDs for the row by hybrid by interval means
dat$row_id <- interaction(dat$block, dat$tree_row)
dat$row_hybrid <- interaction(dat$row_id, dat$hybrid_code)
dat$row_interval <- interaction(dat$row_id, dat$interval)

fit <- lmerTest::lmer(
  PAIC ~ block + treatment_code * hybrid_code * interval +
    (1 | row_id) + (1 | row_hybrid) + (1 | row_interval),
  data = dat, REML = TRUE,
  contrasts = list(block = "contr.sum",
    treatment_code = "contr.sum", hybrid_code = "contr.sum",
    interval = "contr.sum")
)
anova(fit, type = 3, ddf = "Satterthwaite")
emmeans(fit, ~ treatment_code | interval, weights = "equal",
        lmer.df = "satterthwaite")
# Holm adjustment spans all three pairs and all intervals
# within a phase; see treatment_presentation.R.'''
p=insert(anchor,code)
p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(9)
p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_together=True
for run in p.runs: run.font.name="Consolas";run.font.size=Pt(9)
insert(anchor,"The asterisks in the formula expand all main effects and interactions. The terms beginning with 1 | are random intercepts: row for persistent differences among treatment units, row by hybrid for persistent local hybrid differences, and row by interval for fluctuations shared by hybrids in a row. Block is a fixed adjustment; it is distinct from these random terms.")
insert(anchor,"Tests of the block effect","Heading 1")
t=d.add_table(rows=1,cols=4)
t.autofit=False
widths=[2.6,1.1,1.4,1.4]
for i,(title,w)in enumerate(zip(["Model","F statistic","Degrees of freedom","p value"],widths)):
 t.rows[0].cells[i].text=title;t.columns[i].width=Inches(w)
for phase,label in [("active","Active cropping"),("post","After cropping ceased"),("active_with_boundary","Active including boundary")]:
 with (root/f"results/tables/treatment_PAIC_{phase}_tests.csv").open(encoding="utf-8-sig",newline="")as f:
  r=next(r for r in csv.DictReader(f)if r["term_code"]=="block")
 vals=[label,f'{float(r["F value"]):.3f}',f'{float(r["NumDF"]):.0f} and {float(r["DenDF"]):.2f}',f'{float(r["Pr(>F)"]):.3f}']
 for cell,value in zip(t.add_row().cells,vals):cell.text=value
for ri,row in enumerate(t.rows):
 trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement("w:cantSplit"))
 if ri==0:trpr.append(OxmlElement("w:tblHeader"))
 for ci,cell in enumerate(row.cells):
  cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement("w:tcBorders")
  for edge in ["top","left","bottom","right"]:
   b=OxmlElement("w:"+edge);b.set(qn("w:val"),"single");b.set(qn("w:sz"),"5");b.set(qn("w:color"),"D9D9D9");borders.append(b)
  tcpr.append(borders)
  shade=OxmlElement("w:shd");shade.set(qn("w:fill"),"24435A"if ri==0 else("F2F5F7"if ri%2==0 else"FFFFFF"));tcpr.append(shade)
  margin=OxmlElement("w:tcMar")
  for edge in ["top","left","bottom","right"]:
   m=OxmlElement("w:"+edge);m.set(qn("w:w"),"80");m.set(qn("w:type"),"dxa");margin.append(m)
  tcpr.append(margin)
  for p in cell.paragraphs:
   p.paragraph_format.space_before=Pt(3);p.paragraph_format.space_after=Pt(3)
   if ci>0:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
   for run in p.runs:
    run.font.size=Pt(9);run.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255)if ri==0 else RGBColor(0,0,0)
anchor._p.addprevious(t._tbl)
insert(anchor,"These are Type III F-tests with Satterthwaite denominator degrees of freedom from the fitted models. They provide no clear evidence of mean block differences, but do not demonstrate that blocks are identical.")
insert(anchor,"We retain block because it is part of the allocation design. Its purpose is to adjust treatment comparisons for between-block differences, not to identify a substantive effect for the paper. Removing it only when its test is nonsignificant makes the inferential model depend on a preliminary test. The three blocks remain the treatment replicates; retaining block does not turn individual trees into extra independent replicates.")
p=insert(anchor,"Methodological reference: Frey, Hartung, Ogutu and Piepho (2024), Analyze as randomized: Why dropping block effects in designed experiments is a bad idea. Agronomy Journal 116, 1371-1381. https://doi.org/10.1002/agj2.21570")
for run in p.runs:run.font.size=Pt(8.5)
p=insert(anchor,"");p.add_run().add_break(WD_BREAK.PAGE)
post=next(p for p in d.paragraphs if p.text=="Discussion of the post cropping legacy")
insert(post,"Growth during active cropping","Heading 1")
p=insert(post,"");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
inline=p.add_run().add_picture(str(root/"results/figures/treatment_PAIC_active_means.png"),width=Inches(6.1))
inline._inline.docPr.set("descr","Model adjusted PAIC during active cropping averaged across hybrids with phase wide Holm significance markers")
insert(post,"Figure 1 Model-adjusted PAIC during active cropping, averaged equally across hybrids. Asterisks indicate continuous cropping versus continuous fallow differences; plus signs indicate continuous cropping versus alternate cropping differences, after Holm correction across all 21 pair-by-interval tests. Error bars are pointwise 95 percent confidence intervals.","Caption")
insert(post,"The x-axis gives measurement years, not single annual growth years. In particular, the 2001 point averages growth years 1998 to 2000. The three-year measurement gap prevents us from identifying the annual trajectory within that interval; it does not itself explain the low annualised mean. Nonsignificant contrasts at the late visits are not evidence of equal growth.")
p=insert(post,"");p.add_run().add_break(WD_BREAK.PAGE)
for p in d.paragraphs:
 if p.style.name=="Subtitle" and "Updated 20 October 2026"in p.text:p.text=p.text.replace("Updated 20 October 2026","Updated 22 October 2026")
 if p.style.name=="Caption" and p.text.startswith("Figure 1 Model-adjusted PAIC averaged across hybrids after cessation"):
  p.text=p.text.replace("Figure 1","Figure 2",1)
try:
 d.save(path)
except PermissionError:
 path=path.with_name(path.stem + "_updated_20261022.docx")
 d.save(path)
print(path)



