from pathlib import Path
from shutil import copy2
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path.cwd()/"poplar_paper"
path=root/"manuscript/analysis_notes/03_cropping_treatment_and_legacy.docx"
backup=path.parent/"_history"/"03_cropping_treatment_and_legacy_before_2026_10_22.docx"
backup.parent.mkdir(exist_ok=True)
if not backup.exists(): copy2(path,backup)
d=Document(path)
if any(p.text=="Model specification in pseudo R" for p in d.paragraphs):
    raise RuntimeError("Model section already exists; inspect before editing again")
for p in d.paragraphs:
    if p.style.name=="Subtitle":
        p.text=p.text.replace("20 October 2026","22 October 2026")
    if p.style.name=="Caption" and p.text.startswith("Figure 1"):
        p.text=p.text.replace("Figure 1","Figure 2",1)
anchor=next(p for p in d.paragraphs if p.text=="Main results")
def insert(text,style=None):
    return anchor.insert_paragraph_before(text,style)
insert("Model specification in pseudo R","Heading 1")
insert("Fit the same model separately to the active, post-cropping and active-with-boundary datasets. The response is the mean PAIC of available trees in each block by treatment row by hybrid by observed interval. Block, treatment, hybrid and interval are categorical factors. Missing PAIC records are excluded before aggregation.")
code="""# One record per treatment row, hybrid and growth interval
dat <- row_hybrid_interval_means_for_selected_phase
dat$row_id <- interaction(dat$block, dat$tree_row)
dat$row_hybrid <- interaction(dat$row_id, dat$hybrid_code)
dat$row_interval <- interaction(dat$row_id, dat$interval)

fit <- lmerTest::lmer(
  PAIC ~ block + treatment_code * hybrid_code * interval +
    (1 | row_id) + (1 | row_hybrid) + (1 | row_interval),
  data = dat, REML = TRUE,
  contrasts = sum_to_zero_contrasts_for_fixed_factors
)
anova(fit, type = 3, ddf = "Satterthwaite")"""
p=insert(code)
p.paragraph_format.space_before=Pt(3)
p.paragraph_format.space_after=Pt(9)
p.paragraph_format.line_spacing=1.0
p.paragraph_format.keep_together=True
for r in p.runs:
    r.font.name="Consolas";r.font.size=Pt(9)
insert("The asterisk expands to all main effects and two-way and three-way interactions among treatment, hybrid and interval. The block term is additive. The random intercepts account for persistent differences among treatment rows, persistent local hybrid differences within rows, and fluctuations shared by hybrids in the same row and interval. There is no separate tree random intercept because trees have already been averaged within each row and hybrid.")
insert("Treatment contrasts are obtained from estimated marginal means, averaging equally over hybrids. Overall contrasts additionally average equally over observed intervals. The model is unchanged by displaying treatments without separate hybrid panels.")
insert("Block tests and why block is retained","Heading 1")
insert("The fixed block term was nonsignificant in both primary models: active cropping F(2, approximately 4) = 0.798, p = 0.511; post cropping F(2, approximately 4) = 0.481, p = 0.650. These are Type III tests with Satterthwaite denominator degrees of freedom. They indicate insufficient evidence of an overall block difference after adjustment for the other model terms; with only three blocks, they do not establish that block effects are zero.")
insert("We retain block because treatments were allocated within blocks. It adjusts for baseline spatial differences among those blocks and respects the design when comparing treatments. Its inclusion is specified by the experiment, rather than selected by a preliminary significance test. Removing it solely because p exceeds 0.05 would make the analysis depend on a low-power screening test. The three blocks still provide treatment replication, while neighbouring trees remain subsamples; a nonsignificant block term does not make all trees independent replicates.")
insert("Block test sources: results/tables/treatment_PAIC_active_tests.csv and results/tables/treatment_PAIC_post_tests.csv.")
insert("").add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)
insert("Growth during active cropping","Heading 1")
p=insert("");p.alignment=WD_ALIGN_PARAGRAPH.CENTER
shape=p.add_run().add_picture(str(root/"results/figures/treatment_PAIC_active_means.png"),width=Inches(6.1))
shape._inline.docPr.set("descr","Active cropping PAIC by treatment averaged across hybrids with pointwise confidence intervals and phase-wide Holm significance markers.")
insert("Figure 1 Model-adjusted PAIC during active cropping, averaged equally across hybrids. An asterisk marks continuous cropping versus continuous fallow; a plus marks continuous cropping versus alternate cropping. Markers use Holm adjustment across all 21 treatment-pair by interval tests in this phase. Error bars are pointwise 95 percent confidence intervals.","Caption")
insert("Continuous cropping has lower estimated PAIC than continuous fallow throughout the observed active phase. After phase-wide adjustment, this comparison is significant at the 1995 to 1998 visits; continuous cropping also differs from alternate cropping at the 1997 and 1998 visits. Absence of a marker at other visits does not establish equal growth.")
insert("The low point at the 2001 visit is an annualised mean over growth years 1998 to 2000. The missing field visits limit annual resolution but do not themselves explain the low PAIC. Treatment differences narrow in this interval; the possible weather explanation remains untested with interval-aligned predictors.")
insert("Figure source: results/figures/treatment_PAIC_active_means.png. Interval inference: results/tables/treatment_PAIC_active_interval_inference.csv.")
insert("").add_run().add_break(__import__("docx").enum.text.WD_BREAK.PAGE)
d.save(path)
print(path)

