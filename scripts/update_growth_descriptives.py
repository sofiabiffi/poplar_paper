from pathlib import Path
from datetime import datetime
import csv,shutil
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
root=Path(__file__).resolve().parents[1]
p=root/'manuscript/analysis_notes/01_tree_growth.docx'
b=root/'manuscript/analysis_notes/_history'/('01_tree_growth_before_descriptives_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx')
b.parent.mkdir(exist_ok=True);shutil.copy2(p,b)
d=Document(p)
def after(anchor,text,style=None):
 node=OxmlElement('w:p');anchor._p.addnext(node);q=Paragraph(node,anchor._parent);q.text=text
 if style:q.style=style
 return q
for q in d.paragraphs:
 if q.style.name=='Subtitle': q.text='Analysis notes for paper drafting   |   Updated 23 October 2026'
 if q.text.startswith('The table gives 2023 medians.'):
  q.text=('PAIC varied markedly among visits, whereas MAIC showed smoother changes in accumulation since planting (Figure 1). The highest observed median PAIC was 3.02 for Beaupre and 2.41 for Gibecq, both in 1998; Robusta peaked at 2.43 in 2008 and Trichobel at 3.72 in 2021 (all increments in Mg C ha\u207b\u00b9 yr\u207b\u00b9). These are maxima among observed visits, not reconstructed annual maxima during measurement gaps.')
  q=after(q,'Median MAIC for Beaupre, Gibecq and Robusta reached its observed maximum in 2009 (1.52, 1.29 and 1.41, respectively), compared with 1.29, 1.09 and 1.23 in 2023. Trichobel reached its highest median MAIC at the final visit, at 2.13. By 2023, Trichobel had the highest median PAIC (2.33), compared with 0.70 for Beaupre, 0.42 for Gibecq and 0.86 for Robusta (Table 1).')
  q=after(q,'Block-adjusted comparisons of 2023 means supported higher PAIC and MAIC in Trichobel than in every other hybrid (all Holm-adjusted p < 0.001). Beaupre also had higher MAIC than Gibecq (adjusted p = 0.042); the remaining pairwise differences were not significant after correction. These tests concern the final visit and do not establish significant differences at every earlier visit.')
  after(q,'Table 1 Aboveground carbon increment and stock at the 2023 visit. Values are tree-level medians [25th–75th percentiles]. PAIC and MAIC are in Mg C ha\u207b\u00b9 yr\u207b\u00b9; stock is in Mg C ha\u207b\u00b9. PAIC covers the 2022–2023 measurement interval and MAIC covers 31 years since planting.','Caption')
 if q.text.startswith('Use PAIC to discuss'):
  q.text='Use PAIC to discuss recent growth and MAIC or stock to discuss accumulated carbon. The 2023 tests use only three block replicates, so smaller hybrid contrasts remain uncertain. A row-level sensitivity analysis also supports all Trichobel contrasts and the Beaupre–Gibecq MAIC contrast, but additionally identifies Robusta–Gibecq differences; those additional comparisons are therefore sensitive to the replication assumption. The tests do not incorporate allometric uncertainty or establish differences across the full time series. The supplied Robusta 2023 MAIC is retained even though one stock value is missing; metric-specific sample sizes are reported.'
# Add methods without changing existing text.
anchor=next(q for q in d.paragraphs if q.text.startswith('PAIC is the change'))
after(anchor,'For the final-visit hybrid comparison, tree values were first averaged within each row and hybrid (five trees), then equally across the three treatment rows within each block. Separate additive linear models for PAIC and MAIC included block and hybrid, using 12 block-by-hybrid means (three blocks per hybrid; residual df = 6). All six pairwise hybrid comparisons for each outcome were tested, with Holm correction across the combined family of 12 comparisons. Tests compare means; the table retains medians and quartiles to match the figure. The Gaussian models assume independent blocks and common residual variance; three blocks provide limited scope to assess those assumptions.')
rows=list(csv.DictReader((root/'results/tables/growth_2023_descriptive.csv').open()))
lookup={(r['hybrid'],r['metric']):r for r in rows}
letters={'PAIC':['b','b','b','a'],'MAIC':['b','c','bc','a']}
t=d.tables[0]
for c,text in zip(t.rows[0].cells,['Hybrid','PAIC','MAIC','Carbon stock']):c.text=text
for i,h in enumerate(['Beaupre','Gibecq','Robusta','Trichobel']):
 t.rows[i+1].cells[0].text=h
 for j,m in enumerate(['PAIC','MAIC','carbon_stock_mg_ha'],1):
  r=lookup[h,m];cell=t.rows[i+1].cells[j];cell.text=f"{float(r['median']):.2f} [{float(r['q25']):.2f}–{float(r['q75']):.2f}]"
  if m in letters:
   run=cell.paragraphs[0].add_run(letters[m][i]);run.font.superscript=True
  for para in cell.paragraphs:
   for run in para.runs:run.font.size=Pt(9)
node=OxmlElement('w:p');t._tbl.addnext(node);note=Paragraph(node,t._parent)
note.text='n = 45 trees per hybrid for PAIC and MAIC; stock n = 44 for Robusta and 45 for the other hybrids. Letters within each increment column summarise comparisons of block-adjusted means: hybrids without a shared letter differ at Holm-adjusted p < 0.05 across 12 tests. Shared letters indicate insufficient evidence of a difference. Stock is descriptive and has not been tested separately.'
for r in note.runs:r.font.size=Pt(9)
anchor=next(q for q in d.paragraphs if q.text=='results/checks/response_comparison.csv')
for name in ['scripts/calc-descriptives.r','results/tables/growth_2023_descriptive.csv','results/tables/growth_2023_hybrid_contrasts.csv','results/tables/growth_2023_row_sensitivity.csv']:
 anchor=after(anchor,name)
d.save(p)
print(p)
