from pathlib import Path
from datetime import datetime
import shutil
from docx import Document
root=Path(__file__).resolve().parents[1]
script=root/'scripts/trace_soil_provenance.py'
s=script.read_text()
s=s.replace('provisional_tree_row=', 'tree_row=').replace('provisional_block=', 'block=').replace('provisional_former_treatment=', 'former_treatment=')
s=s.replace('Provisional row numbering requires geographic anchor','Confirmed from coordinates, original design and user orientation 2026-10-22')
script.write_text(s)
readme=root/'results/checks/soil_provenance/README.md'
s=readme.read_text()
s=s.replace('## Provisional former treatment mapping','## Confirmed former treatment mapping')
s=s.replace("The coordinate bands are clear. The geographic connection of original tree-row numbers to these bands needs confirmation. Assuming original rows increase from west to east, consistent with the screenshot and colleague's account:","The user confirmed on 22 October 2026 that tree rows run approximately north to south, with row 6 furthest west. Together with the numbered original design and colleague's row/block/treatment account, this anchors the western, middle and eastern sampled bands to rows 6, 10 and 14 respectively:")
s=s.replace('That file or a georeferenced original row map may resolve numbering.','Row ordering is now anchored by the user\'s site-orientation confirmation.')
s=s.replace('- The north arrow and numbered treatment rows in the supplied screenshot support the provisional ordering but do not replace a geographic anchor. The historic screenshot covers Leeds and Silsoe; confirming the site\'s actual orientation remains appropriate.','- Geographic numbering was confirmed by the user on 22 October 2026; exact hybrid segment boundaries and year-specific cropping sides remain unassigned.')
s=s.replace('including provisional management columns and explicit status flags','including confirmed former-management columns and explicit status flags')
readme.write_text(s)
path=root/'manuscript/analysis_notes/04_soil_carbon_stocks.docx'
(path.parent/'_history').mkdir(parents=True, exist_ok=True); shutil.copy2(path,path.parent/'_history'/('04_soil_carbon_stocks_before_orientation_confirmation_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx'))
d=Document(path)
for p in d.paragraphs:
    text=p.text
    text=text.replace("The three agroforestry sampling bands are consistent with the colleague's description of rows 6, 10 and 14, but their geographic numbering remains provisional until an original site anchor is confirmed.","The user confirmed on 22 October 2026 that rows run approximately north to south, with row 6 furthest west. Together with the original design and colleague's account, this confirms the western, middle and eastern sampled bands as rows 6, 10 and 14 respectively.")
    text=text.replace("Geographic bands are inferred from coordinates; treatment row numbers use the colleague's account and the screenshot's west-to-east ordering provisionally.","Geographic bands are inferred from coordinates. Treatment row numbers use the colleague's account, the original numbered design and the user's confirmation on 22 October 2026 that row 6 is furthest west and rows run approximately north to south.")
    text=text.replace('The row and treatment columns in that file are explicitly provisional.','The row, block and former-treatment columns are confirmed using the supplied site orientation.')
    text=text.replace('Confirm that the western, middle and eastern sampled bands are rows 6, 10 and 14.','The western, middle and eastern sampled bands are now confirmed as rows 6, 10 and 14 using the user\'s site-orientation clarification.')
    text=text.replace('that file is not in the supplied sources and may resolve the anchor.','that file is not in the supplied sources and may provide finer sampling-layout details.')
    if text!=p.text:p.text=text
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            if c.text=='Provisional row and block':c.text='Confirmed row and block'
d.save(path)
print('Updated crosswalk builder, audit explanation and Word note; prior note preserved.')
