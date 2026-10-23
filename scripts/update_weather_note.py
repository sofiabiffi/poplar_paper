from pathlib import Path
from datetime import datetime
import shutil, csv
from docx import Document
from docx.shared import Inches

root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'{path.stem}_before_synchrony_{datetime.now():%Y%m%d_%H%M%S}.docx')
d=Document(path)
marker='Shared fluctuations and weather model update'
if any(p.text==marker for p in d.paragraphs):
    raise RuntimeError('Update already present; preserve revisions and edit deliberately.')
# Preserve all existing paragraphs, tables, formatting and manual revisions.
d.add_page_break()
d.add_heading(marker,1)
d.add_paragraph('Update 20 October 2026. PAIC fluctuations are strongly shared across hybrids after broad trajectories are removed. The simple temperature and rainfall model did not improve held-out prediction. Earlier sections retain the original seasonal PCA record; the current decisions and results follow below.')
d.add_heading('Methods',2)
d.add_paragraph('For the synchrony analysis, use the same tree-level hybrid medians as the PAIC and MAIC descriptive figure, restricted to 26 measurement years with positive PAIC for all four hybrids. Fit a separate quadratic age trend to log PAIC for each hybrid, without a shared year effect. Divide each hybrid residual series by its own standard deviation to compare departures. Calculate descriptive Pearson and Spearman correlations; repeat with linear detrending and the legacy weather-response summaries. No missing years are imputed. The last ten observations, 2014 to 2023, form an exploratory descriptive subset, not an estimated change point.')
d.add_paragraph('The separate weather comparison uses 104 legacy hybrid summaries over 26 years. Gamma log-link mixed models include hybrid-specific quadratic age trends and a random measurement-year intercept. Compare the baseline with shared April to September mean temperature and rainfall effects. Leave out all four hybrids from one year together, refit, and predict the unseen year by integrating over the random intercept. This assesses interpolation, not future forecasting.')
d.add_heading('Calculated results',2)
with (root/'results/tables/paic_synchrony_correlations.csv').open(newline='') as f: rows=list(csv.DictReader(f))
r=[float(x['pearson_r']) for x in rows if x['source']=='Tree medians' and x['method']=='quadratic_residual' and 'T' not in (x['hybrid_code_1'],x['hybrid_code_2'])]
d.add_paragraph(f'After quadratic detrending, correlations among Beaupre, Gibecq and Robusta range from {min(r):.2f} to {max(r):.2f}. Trichobel correlations are 0.73 with Beaupre, 0.72 with Gibecq and 0.89 with Robusta. Linear detrending also shows strong positive synchrony. Trichobel correlation with the mean standardised departure of the other hybrids is 0.95 in 1994 to 2013 and 0.61 in 2014 to 2023. Its median raw PAIC ratio to the geometric mean of the others is 1.50 and 2.73, respectively. These comparisons have no inferential uncertainty assessment.')
d.add_paragraph('Held-out RMSE is 1.696 for the age baseline and 1.703 after adding temperature and rainfall. Mean Gamma deviance is 0.607 and 0.612, respectively. Both weather coefficient intervals include zero. Hybrid-specific weather slopes and winter extensions were not fitted because the shared-weather model did not improve either metric. Final full and fold fits have no recorded convergence or singularity warnings.')
d.add_heading('Discussion hypotheses',2)
d.add_paragraph('Shared departures are consistent with a common environmental driver but do not identify weather as its cause. Trichobel maintains higher relative growth in later observations, especially when other hybrids have low PAIC. This is not proof of greater drought tolerance or a distinct weather response; ratios can increase because their denominator declines.')
d.add_heading('Open decisions and confirmed interval convention',2)
d.add_paragraph('Retain weather from year t minus 1, including multi-year PAIC intervals, as confirmed by the user. Annualising PAIC accounts for interval duration, but growth timing within unsampled years remains unknown. State this in Methods and Discussion. Revisit age-curve sensitivity, temporal dependence, weather exposure summaries and coverage, the Robusta response discrepancy, and validation. The PCA retention caveat remains open: weather variance is not PAIC relevance. See references/WEATHER_PCA_REVIEW.md for the saved Wicklin blog and methodological reference.')
d.add_page_break()
d.add_heading('Shared PAIC fluctuations',1)
d.add_picture(str(root/'results/figures/paic_shared_fluctuations_preview.png'),width=Inches(6.65))
d.add_paragraph('Descriptive comparisons use tree medians. Correlations do not establish causation, and quadratic trends are provisional. The Trichobel ratio compares observed PAIC with the geometric mean of Beaupre, Gibecq and Robusta.',style='Caption')
d.add_heading('Reproducible sources for this update',2)
d.add_paragraph('analysis/07_weather_temperature_rainfall.R and analysis/08_check_weather_temperature_rainfall.R; analysis/09_paic_shared_fluctuations.R; results/tables/paic_synchrony_*.csv; results/tables/weather_temperature_rainfall_*.csv; references/WEATHER_MODEL_NOTES.md and references/WEATHER_PCA_REVIEW.md.')
d.save(path)
print(path)
