from pathlib import Path
from datetime import datetime
import csv,shutil
from docx import Document
from docx.shared import Inches
from docx.text.paragraph import Paragraph

root=Path(r'C:/Users/sofia/Documents/poplar_paper')
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'02_weather_and_PAIC_before_model_se_{datetime.now():%Y%m%d_%H%M%S}.docx')
d=Document(path)
start=next(p for p in d.paragraphs if p.text in ('Bootstrap uncertainty','Model standard errors'))
end=next(p for p in d.paragraphs if p.text=='Calculated weather associations by hybrid')
node=start._p
while node is not end._p:
    nxt=node.getnext();node.getparent().remove(node);node=nxt
p=end.insert_paragraph_before('Model standard errors');p.style='Heading 1'
end.insert_paragraph_before('Estimates are presented with ±1 model-based standard error (SE). For the absolute PAIC differences, SEs are calculated by the delta method using the full coefficient covariance matrix. This carries uncertainty from the fitted model onto the carbon-increment scale. The whiskers show ±1 SE, not 95% confidence intervals, and should not be read as a 5% significance test.')
end.insert_paragraph_before('Suggested Methods wording: “Hybrid-specific Gamma models with a log link included summer rainfall, summer temperature, age and age squared. Absolute mean PAIC contrasts were averaged over the observed ages and reported with model-based standard errors calculated by the delta method. Coefficient tests used the conventional model t tests. Standard errors and tests assume independent observations and are conditional on the selected model.”')

def replace(prefix,text):
    ps=[p for p in d.paragraphs if p.text.startswith(prefix)]
    assert len(ps)==1,(prefix,len(ps));ps[0].text=text;return ps[0]

replace('The same quadratic age specification', 'The same quadratic age specification is used in all eight hypothesis models (32 fits on the common 24 years). The combined summer-weather model ranks first in each hybrid and supplies the final associations. Reporting uses all 26 complete-summer years. Its inputs and estimates were verified against the earlier fit. Model SEs and coefficient tests describe this selected model; they do not include model-selection uncertainty.')
caption=replace('Figure 2.', 'Figure 2. Separate Gamma log-link models include linear and squared age, June–August rainfall and mean temperature for each hybrid, using 26 measurement years. Rainfall is contrasted from 122.44 to 222.44 mm at 15.81 °C; temperature is contrasted from 15.31 to 16.31 °C at 172.44 mm rainfall. Differences are averaged equally over the 26 observed ages and expressed in Mg C ha⁻¹ yr⁻¹. Points show absolute mean PAIC differences; whiskers show ±1 model-based SE calculated by the delta method, not 95% confidence intervals. The dashed line marks zero. Weather describes the preceding growth year, including annualised multi-year intervals.')
pic=Paragraph(caption._p.getprevious(),caption._parent)
assert pic._p.xpath('.//w:drawing')
for child in list(pic._p):
    if child.tag.endswith('}r'):pic._p.remove(child)
pic.add_run().add_picture(str(root/'results/figures/final_weather_hybrid_effects.png'),width=Inches(6.35))
table_caption=replace('Table 1.', 'Table 1. Absolute mean PAIC differences ±1 model-based SE (Mg C ha⁻¹ yr⁻¹). Contrasts are centred on study-average weather and averaged over the same observed ages. SEs are calculated from the full model coefficient covariance by the delta method.')
node=table_caption._p.getprevious()
while node.tag.split('}')[-1]!='tbl':node=node.getprevious()
from docx.table import Table
table=Table(node,d._body)
with (root/'results/tables/final_weather_absolute_effects.csv').open(encoding='utf-8-sig') as f:effects=list(csv.DictReader(f))
for row in table.rows[1:]:
    h=row.cells[0].text
    for col,pred in [(1,'Rainfall'),(2,'Temperature')]:
        e=next(e for e in effects if e['hybrid']==h and e['predictor']==pred)
        cell=row.cells[col];para=cell.paragraphs[0]
        text=f"{float(e['difference']):.2f} ± {float(e['standard_error']):.2f}"
        if para.runs:
            para.runs[0].text=text
            for r in para.runs[1:]:r.text=''
        else:para.add_run(text)
replace('Rainfall estimates are positive', 'Rainfall estimates are positive in all hybrids. Trichobel has the largest absolute estimate, 0.77 ± 0.27 Mg C ha⁻¹ yr⁻¹ per 100 mm, followed by Beaupre (0.69 ± 0.22), Gibecq (0.66 ± 0.23) and Robusta (0.56 ± 0.19). Conventional rainfall-coefficient p-values are 0.0020, 0.0046, 0.0052 and 0.0107 for Beaupre, Gibecq, Robusta and Trichobel, respectively.')
replace('Temperature estimates are also positive', 'Temperature estimates are also positive: Beaupre 0.44 ± 0.18, Gibecq 0.39 ± 0.19, Robusta 0.53 ± 0.17 and Trichobel 0.58 ± 0.24 Mg C ha⁻¹ yr⁻¹ per 1 °C. Conventional temperature-coefficient p-values are 0.0152, 0.0396, 0.0026 and 0.0222, respectively. Both weather coefficients have p < 0.05 in each hybrid under these tests. Separate within-hybrid tests do not test whether the hybrid responses differ.')
replace('Trichobel has the largest estimated absolute association', 'Trichobel has the largest estimated absolute association with both rainfall and temperature. This is a descriptive ranking of point estimates; the separate model tests assess each hybrid’s association with weather, not differences between hybrids. Appropriate wording is: “Trichobel showed the largest estimated absolute PAIC association with summer rainfall and temperature.” The whiskers describe uncertainty in mean associations, not variation in physiological sensitivity between years.')
replace('In the 24-year summer models', 'In the 26-year reporting summer models, consecutive-calendar-year residual correlations are approximately 0.35, 0.56, 0.29 and 0.00 for Beaupre, Gibecq, Robusta and Trichobel. The models assume independent observations; remaining temporal dependence may make conventional SEs and p-values optimistic. Results are therefore presented as exploratory associations.')
for p in list(d.paragraphs):
    if p.text.startswith(('All 4,000 bootstrap model fits','Residual, Gamma quantile, variance-pattern')):
        p._p.getparent().remove(p._p)
# Preserve a concise diagnostic account, separate from biological interpretation.
p=next(p for p in d.paragraphs if p.text.startswith('In the 26-year reporting'))
new=d.add_paragraph('Residual, Gamma quantile, variance-pattern and influence plots were inspected. There is some residual temporal structure, variation in residual spread and influential observations; year-omission checks are retained. Diagnostic plots and summaries are saved with the analysis. These checks do not establish that all model assumptions hold.')
p._p.addnext(new._p)
assert not any('bootstrap' in p.text.lower() for p in d.paragraphs)
d.save(path)
description=caption.text
(root/'results/figure_descriptions/final_weather_hybrid_effects.txt').write_text(description+'\n',encoding='utf-8')
print(path)
