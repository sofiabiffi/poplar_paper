"""Build the current soil note and retain the complete manual provenance section."""
from pathlib import Path
from copy import deepcopy
from datetime import datetime
from docx import Document
from lxml import etree
import shutil,csv,json,hashlib
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/04_soil_carbon_stocks.docx'
heading='Soil sample provenance and former cropping treatments'
old=Document(path); protected=[]; found=False
for element in old.element.body:
 text=''.join(element.xpath('.//w:t/text()'))
 if text.strip()==heading:found=True
 if found and not element.tag.endswith('}sectPr'):protected.append(deepcopy(element))
if not protected:raise RuntimeError('Required provenance section not found; refusing to overwrite the note')
def fingerprint(elements):
 return hashlib.sha256(b''.join(etree.tostring(e,method='c14n',exclusive=True) for e in elements)).hexdigest()
protected_hash=fingerprint(protected)
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/('04_soil_carbon_stocks_before_CSV_analysis_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx'))
ns={'__file__':str(root/'scripts/build_analysis_notes.py')}
exec((root/'scripts/build_analysis_notes.py').read_text().split('# Growth')[0],ns)
make,h,p,table,picture,save=[ns[x] for x in ('make','h','p','table','picture','save')]
rows=lambda name:list(csv.DictReader((root/'results/tables'/name).open(encoding='utf-8-sig',newline='')))
ref=rows('soc_ESM_arable_reference_masses.csv'); audit=rows('soc_ESM_extrapolation_audit.csv')
fixed=rows('soc_profile_summary.csv'); esm=rows('soc_ESM_total_summary.csv');pairs=rows('soc_ESM_layer_dunn_BH.csv')
source=list(csv.DictReader((root/'data/raw/study_site_soil_2024.csv').open(encoding='utf-8-sig',newline='')))
n_layers=len(source);n_profiles=len(set(r['profile_id'] for r in source));n_complete=sum(int(r['n_profiles']) for r in fixed)
d=make('Soil organic carbon stocks and equivalent soil mass correction','This analysis quantifies soil organic carbon stocks, bulk density and SOC content in 2024 for the arable reference, alleys and tree rows. Arable reference stocks remain fixed-depth observations. Linear equivalent soil mass correction evaluates alley and tree-row stocks at the arable mean soil masses. Tabular total stocks support subsequent area scaling; reference-mass intervals support depth-related interpretation.')
d.paragraphs[1].text='Analysis notes for paper drafting   |   Updated '+datetime.now().strftime('%d %B %Y')
h(d,'Methods')
p(d,f'The measurement input is data/raw/study_site_soil_2024.csv, containing {n_layers} layers from {n_profiles} permanent company core IDs and {n_complete} complete two-layer profiles. Source values retain full precision. Arable reference, Alley and Tree row identify land cover; former cropping treatments are separate metadata. All profiles were sampled in 2024. Company sample IDs identify layers; core IDs identify profiles. The provenance section below retains the sample-origin record.')
p(d,'Layer soil mass (Mg soil per hectare) equals fine dry bulk density (g/cm3) times layer thickness (cm) times 100. Layer SOC stock (Mg C per hectare) equals soil mass times SOC content (%) divided by 100. Fine BD and SOC content are used directly; no additional organic-matter subtraction or coarse-fragment adjustment is applied. Fixed-depth 0 to 30 cm stocks sum both observed layers and exclude incomplete profiles. Company-supplied stock values are retained as source fields rather than substituted for these calculations.')
p(d,f'The upper reference mass is {float(ref[0]["reference_cumulative_mass_mg_ha"]):.2f} Mg soil per hectare, the mean mass of {ref[0]["n_samples"]} arable 0 to 15 cm samples. The mean lower-layer mass is {float(ref[1]["mean_layer_mass_mg_ha"]):.2f} from {ref[1]["n_samples"]} samples. Their sum, {float(ref[1]["reference_cumulative_mass_mg_ha"]):.2f}, is the arable 0 to 30 cm reference mass. Layer means use all available arable observations.')
p(d,'For each core, join the origin and measured cumulative soil-mass and SOC-stock endpoints with straight segments, without splines. Evaluate stocks at the two reference masses using the segment bracketing each target. Where the upper reference exceeds the measured upper-layer mass in a complete profile, the extra mass takes the measured lower-layer SOC concentration. Complete profiles can extend beyond their measured total mass using the final layer slope. The lower reference-mass stock is total minus upper stock. Arable profiles retain their observed fixed-depth 0 to 15 cm and 15 to 30 cm stocks; they are not ESM corrected. Only alley and tree-row profiles are evaluated at the arable mean reference masses.')
p(d,f'The {n_profiles-n_complete} upper-only profiles are attributed to shallow bedrock. Retain an upper-reference estimate by evaluating their upper-layer slope, extending to the upper reference if necessary. This is a mass-standardisation assumption rather than an observation below bedrock. Do not calculate their lower-reference increment or total stock. Reference mass depth labels denote the arable 0 to 15 and 15 to 30 cm intervals, not equal physical depths in every core.')
p(d,'Supplementary properties use complete measured 0 to 30 cm profiles. Within each core, BD is soil-mass weighted: sum(BD squared times thickness) divided by sum(BD times thickness). SOC content is the simple arithmetic mean of the two layer percentages. These measured-depth properties are not ESM transformed. Location summaries report mean, SD, SE, median and quartiles; SE is SD divided by square root of n. Reference-mass uncertainty is not propagated.')
p(d,'Location tests are Kruskal-Wallis followed by two-sided tie-corrected Dunn comparisons. Benjamini-Hochberg adjustment covers three pairs separately for each reference-mass interval, each fixed-depth interval and each supplementary property. Letters use adjusted p < 0.05; shared letters indicate no detected difference. These tests describe sampled-core contrasts and assume core independence. They do not resolve spatial clustering or identify replicated land-use effects.')
d.add_page_break();h(d,'Calculated results')
fixed_by={r['location_code']:r for r in fixed}
table(d,['Land cover','Profiles','Fixed depth','Reported stock','Reported SD'],[[r['location'],r['n_profiles'],f"{float(fixed_by[r['location_code']]['mean']):.2f}",f"{float(r['mean']):.2f}",f"{float(r['sd']):.2f}"] for r in esm],[1.9,.65,1.2,1.15,1.15])
p(d,'Stocks are Mg C per hectare. Fixed depth is measured 0 to 30 cm for all land covers. Reported comparison stocks use fixed-depth arable values and ESM-corrected alley/tree-row values at the arable 0 to 30 cm reference mass. Complete-profile counts apply to both totals. SD describes variation across cores, not uncertainty from independent experimental land-use replication.')
table(d,['Land cover','Extended totals','Excluded totals','Maximum depth cm'],[[r['location'],r['n_total_extrapolated']+' / '+r['n_complete'],r['n_total_excluded'],f"{float(r['max_equivalent_depth_cm']):.1f}"] for r in audit],[1.9,1.4,1.25,1.5])
p(d,'For alley and tree rows, extended total stocks assume that the measured lower-layer SOC concentration continues below 30 cm. Arable profiles are not extended. Upper-only bedrock profiles retain no total. Upper-reference counts are '+', '.join(r['location']+' '+r['n_profiles'] for r in audit)+'.')
interval_name={'Upper reference mass':'0-15','Lower reference mass increment':'15-30'}
table(d,['Reference depth cm','Land-cover comparison','BH adjusted p'],[[interval_name[r['mass_interval']],r['group1']+' versus '+r['group2'],f"{float(r['p_BH']):.3g}"] for r in pairs],[1.15,3.8,1.1])
letters=rows('soc_ESM_layer_letters.csv')
for interval in interval_name:
 vals=[r['letter'] for r in letters if r['depth_interval_cm']==interval]
 p(d,f'Comparison letters at {interval_name[interval]} cm reference mass depth: '+', '.join(vals)+' in arable, alley and tree-row order.')
d.add_page_break();h(d,'SOC stocks by land cover')
picture(d,'soc_ESM_by_mass_layer.png','Figure 1 SOC stocks: arable profiles use measured fixed-depth 0 to 15 and 15 to 30 cm stocks; alley and tree-row profiles use ESM correction to corresponding arable reference masses. Letters compare land covers within each mass interval. Boxes show median and IQR; whiskers use 1.5 times IQR; points are core estimates.',6.5)
h(d,'Supplementary bulk density and SOC content')
props=rows('supplement_soil_properties_0_30.csv');body=[]
for r in fixed:
 v={x['variable']:x for x in props if x['location']==r['location']}
 fmt=lambda x:f"{float(x['mean']):.3f} ({float(x['sd']):.3f}) {x['significance_letter']}"
 body.append([r['location'],r['n_profiles'],fmt(v['BD_mass_weighted_g_cm3']),fmt(v['SOC_content_mean_percent'])])
table(d,['Land cover','Profiles','BD g/cm3','SOC content %'],body,[1.9,.7,1.7,1.75])
p(d,'Mean (SD) across complete measured 0 to 30 cm cores. BD is mass weighted within core; SOC content is an unweighted layer average. Letters compare land covers separately within each property using Dunn tests with BH adjustment, p < 0.05. Shared letters indicate no detected difference. Properties describe measured depth, not the variable physical depth reaching the ESM reference.')
p(d,'Fixed-depth stocks, depth comparisons and the measured 0 to 30 cm figure are available separately as soc_profile_summary, soc_stock_depth_dunn_BH, soc_stock_by_depth and soc_stock_0_30. The primary total comparison table is soc_stock_comparison_total_summary, with stock_basis identifying fixed-depth arable and ESM agroforestry values.')
d.add_page_break();h(d,'Supplementary cumulative coordinate overview')
picture(d,'supplement_soc_ESM_cumulative_profiles.png','Figure S1 Each facet is an alley or tree-row core. Black points are measured cumulative endpoints joined linearly from the origin. Purple dashed segments extend the final measured slope. Dotted vertical lines show arable reference masses. Blue triangles and orange squares represent ESM-corrected 0 to 15 cm and 0 to 30 cm SOC stocks. Asterisks mark upper-only shallow-bedrock cores.',6.5)
p(d,'The standalone vector PDF shows all '+str(sum(int(r['n_profiles']) for r in audit if r['location_code']!='3_arable'))+' alley and tree-row cores. Facet numbers are company core numbers in agroforestry Field 2, not legacy workbook IDs. The figure-profile key links facet labels to permanent core IDs, coordinates and former-treatment metadata.')
h(d,'Discussion hypotheses')
p(d,'Spatial differences in organic inputs and soil conditions could contribute to land-cover contrasts. Baseline differences, field history and spatial clustering remain alternative explanations. ESM intervals can include soil below 30 cm where extrapolation is required. These cross-sectional contrasts do not establish causal former-treatment effects or measured carbon accumulation through time.')
h(d,'Open decisions')
p(d,'Former cropping treatment is confounded with block and sampled row band in this soil subset, as documented in the provenance record. Cores are spatial subsamples. Upscaling requires tree-row and alley area fractions, a comparable arable baseline and compatible area and carbon units with aboveground stocks. A stock contrast requires an explicit baseline and timescale before interpretation as a sequestration rate.')
# Keep the complete requested section, including all its paragraphs and table, verbatim.
d.add_page_break()
for element in protected:d.element.body.insert(len(d.element.body)-1,deepcopy(element))
save(d,'04_soil_carbon_stocks.docx')
check=Document(path);after=[];found=False
for element in check.element.body:
 if ''.join(element.xpath('.//w:t/text()')).strip()==heading:found=True
 if found and not element.tag.endswith('}sectPr'):after.append(element)
assert fingerprint(after)==protected_hash,'Provenance content changed'
(root/'results/checks/soc_note_provenance_preservation.json').write_text(json.dumps({'heading':heading,'xml_elements':len(protected),'before_sha256':protected_hash,'after_sha256':fingerprint(after),'identical':True},indent=2))
print('Preserved provenance section verbatim:',len(protected),'XML elements')
