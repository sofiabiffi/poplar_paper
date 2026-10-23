from pathlib import Path
from datetime import datetime
import shutil,csv
from docx import Document
from docx.shared import Inches
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
history=path.parent/'_history';history.mkdir(exist_ok=True)
d=Document(path)
title='Exploratory weather screen of shared departures'
if any(p.text==title for p in d.paragraphs): raise RuntimeError('Already updated; edit deliberately to preserve manual changes.')
shutil.copy2(path,history/f'{path.stem}_before_weather_screen_{datetime.now():%Y%m%d_%H%M%S}.docx')
d.add_page_break();d.add_heading(title,1)
d.add_paragraph('Update 20 October 2026. The screen suggests modest positive associations with summer rainfall and temperature. No clear dominant weather driver has emerged, and no selected predictor has been validated as a PAIC model.')
d.add_heading('Methods',2)
d.add_paragraph('Average the four within-hybrid standardised log PAIC departures from their separate quadratic age trajectories, giving hybrids equal weight. Screen nine predefined candidates: spring and summer rainfall, sunshine and mean air temperature; spring air and ground frost-day counts; and the longest dry spell during April to September. Spring is March to May and summer is June to August of growth year t minus 1. Dry days have rainfall below 1 mm; the dry-spell measure is the longest consecutive run, with no missing days allowed. Frost means daily minimum temperature below zero. Retain the confirmed t minus 1 alignment for annualised multi-year PAIC intervals.')
d.add_paragraph('Require complete daily coverage for the particular predictor. The primary comparison uses 24 years complete for all nine measures, omitting measurement 2001 because spring temperature lacks one day and 2019 because spring ground frost lacks two days. Predictor-specific available-year results retain 25 or 26 years. Report Pearson and Spearman correlations, the range of Pearson correlations obtained by omitting each year, correlations after adjusting both series for quadratic age on the matched sample, and sensitivity to linear rather than quadratic PAIC detrending. Omission ranges are influence diagnostics, not confidence intervals. Detrending is fitted to all observed years, so these summaries are not predictive validation.')
d.add_heading('Calculated results',2)
with (root/'results/tables/weather_departure_screen_common_years.csv').open(newline='') as f: rows=list(csv.DictReader(f))
for key in ['summer_rainfall','summer_temp','spring_rainfall']:
    r=next(x for x in rows if x['id']==key)
    d.add_paragraph(f"{r['label']}: Pearson r {float(r['pearson_r']):.2f}, Spearman r {float(r['spearman_r']):.2f}; single-year omission range {float(r['jackknife_min']):.2f} to {float(r['jackknife_max']):.2f}. Linear PAIC detrending gives r {float(r['linear_detrending_r']):.2f}; quadratic age adjustment of both matched series gives r {float(r['age_adjusted_r']):.2f}.")
d.add_paragraph('Other Pearson correlations range from approximately -0.05 to 0.17 on the common basis. Summer rainfall has a positive association within each hybrid, r 0.23 to 0.38; summer temperature also has a positive association in each, r 0.28 to 0.44. Predictor-specific available years give r 0.32 for summer rainfall and 0.33 for summer temperature. Dry spell association is weak and positive rather than the hypothesised negative association; it is not evidence of improved growth under drought.')
d.add_heading('Discussion hypotheses',2)
d.add_paragraph('Summer rainfall is a candidate for follow-up because its direction is consistent across hybrids, ranks and omission checks. The summer temperature association is influenced by the earliest observation, 1994, and is weaker using ranks. Both associations may reflect unmeasured conditions, growth-trajectory misspecification or chance among the nine screened candidates. Seasonal aggregation may capture relationships lost in the earlier half-year model, but this has not yet been tested in a validated model. Shared fluctuations remain substantially unexplained.')
d.add_heading('Open decisions',2)
d.add_paragraph('Next assess a small summer-weather model against the age baseline, with all detrending and any predictor selection inside training folds. Keep year groups together and consider temporal blocks. Check sensitivity to establishment-age observations without deleting them from the primary analysis; retain multi-year intervals and disclose their limitation. Do not interpret the screen as an independent validation set or convert individual unadjusted correlations into confirmatory claims. Weather coverage and the Robusta legacy discrepancy remain open.')
d.add_page_break();d.add_heading('Weather screen overview',1)
d.add_picture(str(root/'results/figures/weather_departure_screen_preview.png'),width=Inches(6.65))
d.add_paragraph('Blue points show Pearson correlations; orange triangles show Spearman rank correlations. Blue ranges show sensitivity to omitting one year at a time and are not confidence intervals.',style='Caption')
d.add_heading('Reproducible sources',2)
d.add_paragraph('analysis/10_weather_departure_screen.R; results/tables/weather_departure_screen_common_years.csv; weather_departure_screen_available_years.csv; weather_departure_screen_data.csv; weather_departure_hybrid_correlations.csv. The nine scatterplots are saved separately as results/figures/weather_departure_scatterplots.pdf and PNG.')
d.save(path);print(path)
