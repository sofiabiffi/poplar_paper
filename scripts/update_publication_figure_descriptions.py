"""Refresh embedded figures and keep their explanatory text in Word captions."""
from pathlib import Path
from datetime import datetime
import shutil
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
NOTES = ROOT/'manuscript/analysis_notes'
MAPPING = {
 '01_tree_growth': ['paic_maic_descriptive'],
 '02_weather_and_PAIC': ['final_weather_raw_screen','final_weather_hybrid_effects','final_weather_quadratic_age_shapes'],
 '03_cropping_treatment_and_legacy': ['treatment_PAIC_active_means','treatment_PAIC_post_means'],
 '04_soil_carbon_stocks': ['soc_ESM_by_mass_layer','supplement_soc_ESM_cumulative_profiles'],
 '06_soc_sequestration': ['soc_sequestration_ESM_stocks','soc_sequestration_ESM_rates'],
}
EXTRA = {
 '01_tree_growth': ['paic_shared_fluctuations','trichobel_relative_paic'],
 '03_cropping_treatment_and_legacy': ['treatment_PAIC_active_contrasts','treatment_PAIC_post_contrasts','treatment_PAIC_active_with_boundary_means','treatment_PAIC_active_with_boundary_contrasts'],
 '04_soil_carbon_stocks': ['soc_stock_by_depth','soc_stock_0_30'],
}
DETAIL = {
 'treatment_PAIC_active_means': 'Model-adjusted periodic annual carbon increment (PAIC) during active cropping, with equal weights across hybrids. Symbols and lines distinguish continuous fallow, continuous cropping and alternate cropping; lines connect observed measurement visits. Error bars are pointwise 95% confidence intervals. An asterisk marks continuous cropping versus continuous fallow and a plus marks continuous cropping versus alternate cropping at phase-wide Holm-adjusted p < 0.05 across all 21 treatment-pair by interval tests, including the fallow versus alternate comparison. Absence of a marker does not establish equality. The horizontal axis gives measurement year; PAIC is annualised over the preceding measurement interval, including growth years 1998 to 2000 for the 2001 visit.',
 'treatment_PAIC_post_means': 'Model-adjusted periodic annual carbon increment (PAIC) after cropping ceased, with equal weights across hybrids. Symbols and lines distinguish the three former cropping treatments. Error bars are pointwise 95% confidence intervals. An asterisk marks former continuous cropping versus former continuous fallow and a plus marks former continuous cropping versus former alternate cropping at phase-wide Holm-adjusted p < 0.05 across all 54 treatment-pair by interval tests. The significant fallow contrasts occur at the 2009 and 2010 visits, corresponding to growth years 2008 and 2009. Absence of a marker does not establish equality. The horizontal axis gives measurement year and PAIC is annualised over the preceding measurement interval.',
 'soc_ESM_by_mass_layer': 'SOC stocks for arable reference, alley and tree-row profiles. Arable stocks remain measured fixed-depth 0-15 cm and 15-30 cm values. Alley and tree-row stocks are ESM corrected to the corresponding arable mean reference masses. Boxes show median and IQR; whiskers extend to observations within 1.5 times the IQR. Points are cores. Dunn letters use BH adjustment across three pairs within each panel, p < 0.05. Shared letters indicate no detected difference.',
 'supplement_soc_ESM_cumulative_profiles': 'Each facet shows one alley or tree-row core, labeled by land cover and company core number in Field 2, on shared axes. Black points show measured cumulative soil-mass and SOC-stock endpoints joined linearly from the origin. Purple dashed segments extend the final measured layer using its SOC concentration. Grey dotted lines mark arable mean cumulative reference masses of 1775 and 3733 Mg soil per hectare. Coloured symbols identify ESM-corrected stocks at the upper and total reference masses, as identified in the legend. An asterisk in a facet label indicates a shallow-bedrock core for which only an upper-reference estimate is retained; no lower increment or full-profile estimate is calculated. Reference depth labels describe arable depths rather than actual depths in each core.',
 'trichobel_relative_paic': 'Observed tree-level median PAIC for Trichobel divided by the geometric mean of Beaupre, Gibecq and Robusta at each measurement visit. The dashed line at one indicates equal PAIC. Symbols are observed ratios and lines connect visits without imputing missing years. The ratio is descriptive and does not adjust for age or quantify uncertainty.',
 'soc_sequestration_ESM_stocks': 'SOC stocks in the arable reference, alley fallow and tree rows. Arable stocks are measured at fixed depth (0-30 cm); agroforestry stocks are ESM corrected to the arable 0-30 cm reference mass. Points show complete-profile means and error bars show SD among profiles, not confidence intervals. The dashed line marks the measured fixed-depth arable mean. Groups comprise 47 arable profiles, eight profiles in each selected alley group and 21 tree-row profiles.',
 'paic_maic_descriptive': 'PAIC is periodic annual carbon increment and MAIC is mean annual carbon increment. Symbols and lines identify hybrid medians; shaded ribbons span the 25th to 75th percentiles across trees. The panels use separate vertical scales. PAIC is annualised over the interval between measurements; the 2001 observation therefore averages growth years 1998 to 2000. Lines connect observed visits; missing years are not imputed.',
 'final_weather_raw_screen': 'Cell values and colour show Pearson correlation coefficients on the same 24 complete-weather measurement years for all nine candidates. PAIC is periodic annual carbon increment, paired with weather in the preceding growth year. No age adjustment or response standardisation is applied in this screen. The nine candidates were hypothesis guided; correlations alone did not determine model selection and may reflect shared temporal trends.',
 'final_weather_hybrid_effects': 'Separate Gamma models with a log link include linear and squared age, June to August rainfall and June to August mean temperature for each hybrid, using 26 measurement years. Rainfall is contrasted from 122.44 to 222.44 mm while temperature is held at 15.81 degrees C; temperature is contrasted from 15.31 to 16.31 degrees C while rainfall is held at 172.44 mm. Differences are averaged equally over the same 26 observed ages and expressed in Mg C per hectare per year. Whiskers are pointwise 95% intervals from 1,000 paired circular three-visit block-bootstrap resamples; they describe uncertainty in mean associations, not annual variability. The dashed vertical line marks zero association. Weather is aligned to the preceding growth year, including annualised multi-year growth intervals.',
 'final_weather_quadratic_age_shapes': 'The curves are Gamma log-link fitted means with hybrid-specific linear and squared age terms, holding June to August mean temperature at 15.81 degrees C and rainfall at 172.44 mm. Only the observed age range is shown. Uncertainty bands are omitted. These curves describe conditional fitted shapes, not observed trajectories or independently identified causal effects of age.',
 'paic_shared_fluctuations': 'Descriptive comparison of hybrid PAIC on common positive-response measurement years. Panel A shows observed tree-level hybrid medians as thin lines and symbols and separate quadratic trends fitted to log PAIC as thick lines. Panel B shows log-scale residuals divided by the within-hybrid residual standard deviation; zero denotes the fitted trend. Panel C gives Pearson correlations of those standardised departures. Panel D divides Trichobel PAIC by the geometric mean for Beaupre, Gibecq and Robusta; the dashed line at one denotes equal PAIC. No shared year effect is removed. Lines connect observed visits without imputing missing years. Correlations do not establish weather causation; sensitivity analyses use linear detrending and the legacy response series.',
 'soc_sequestration_ESM_rates': "Annualised SOC accumulation estimated as each complete agroforestry profile's ESM stock minus the mean measured fixed-depth arable stock, divided by the agreed duration (19 years for former Continuous cropping alleys; 32 years for Continuous fallow alleys and tree rows). Points show group means and error bars show SD among profile-level annualised estimates. The shared arable baseline is held fixed; SD is not a confidence interval. Arable zero is the defined reference. These are retrospective stock contrasts, not directly measured annual fluxes.",
}
HISTORY = NOTES/'_history'/('figure_presentation_'+datetime.now().strftime('%Y%m%d_%H%M%S'))
HISTORY.mkdir(parents=True,exist_ok=True)

def description(stem):
    if stem in DETAIL: return DETAIL[stem]
    path = ROOT/'results/figure_descriptions'/f'{stem}.txt'
    text = ' '.join(path.read_text().split()) if path.exists() else ''
    return text

for docstem, figures in MAPPING.items():
    path = NOTES/f'{docstem}.docx'
    shutil.copy2(path,HISTORY/path.name)
    # Required pre-presentation templates are retained independently of backups.
    original = ROOT/'references/note_templates'/path.name
    if not original.exists():
        original = path
    d = Document(original)
    image_paras = [p for p in d.paragraphs if p._p.xpath('.//w:drawing')]
    assert len(image_paras)==len(figures),(docstem,len(image_paras))
    for para, stem in zip(image_paras,figures):
        drawing = para._p.xpath('.//wp:inline')[0]
        width = int(drawing.extent.cx)
        if stem in ('final_weather_raw_screen','final_weather_quadratic_age_shapes','supplement_soc_ESM_cumulative_profiles','paic_maic_descriptive'):
            width=int(width*.94)
            drawing.extent.cx=width
            for ext in para._p.xpath('.//a:xfrm/a:ext'): ext.set('cx',str(width))
        image_path = ROOT/'results/figures'/f'{stem}.png'
        # Keep the existing drawing and relationship; only replace its image bytes.
        rid = para._p.xpath('.//a:blip')[0].get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
        d.part.related_parts[rid]._blob = image_path.read_bytes()
        with Image.open(image_path) as im: height = int(width*im.height/im.width)
        drawing.extent.cy = height
        for ext in para._p.xpath('.//a:xfrm/a:ext'): ext.set('cy',str(height))
        para.paragraph_format.keep_with_next=True
        nxt = para._p.getnext()
        cap = Paragraph(nxt,para._parent)
        assert cap.style.name=='Caption',(stem,cap.text)
        existing = cap.text
        detail = description(stem)
        import re
        prefix=re.match(r'Figure\s+(?:S?\d+)\.?',existing).group()
        # Custom descriptions are complete replacements; extracted notes supplement
        # existing captions so pre-existing interpretation is retained.
        if stem in DETAIL: cap.text=prefix+' '+detail
        elif detail and detail not in existing: cap.text = existing+' '+detail
        drawing.docPr.set('descr',cap.text)
    extras = EXTRA.get(docstem,[])
    if extras and not any(p.text=='Descriptions for additional figure exports' for p in d.paragraphs):
        heading=d.add_heading('Descriptions for additional figure exports',1)
        heading.paragraph_format.page_break_before=True
        for stem in extras:
            d.add_paragraph(stem.replace('_',' '),'Heading 2')
            d.add_paragraph(description(stem))
            d.add_paragraph('Figure file: results/figures/'+stem+'.png')
    d.save(path)
    print(path.name)
