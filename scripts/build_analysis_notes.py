from pathlib import Path
import csv
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "manuscript" / "analysis_notes"
OUT.mkdir(parents=True, exist_ok=True)
def rows(name):
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))
def num(v, digits=2): return f"{float(v):.{digits}f}"
def make(title, summary):
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    sec.left_margin = sec.right_margin = Inches(0.8)
    for name in ["Normal","Title","Subtitle","Heading 1","Heading 2","Caption"]:
        s = d.styles[name]
        s.font.name = "Arial"
        s.font.color.rgb = RGBColor(0,0,0)
    s = d.styles["Normal"]; s.font.size = Pt(10.5)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.08
    d.styles["Title"].font.size = Pt(21)
    d.styles["Title"].paragraph_format.space_after = Pt(9)
    d.styles["Heading 1"].font.size = Pt(13)
    d.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    d.styles["Heading 1"].paragraph_format.space_after = Pt(6)
    d.styles["Caption"].font.size = Pt(9)
    d.add_paragraph(title, "Title")
    d.add_paragraph("Analysis notes for paper drafting   |   Updated 20 October 2026", "Subtitle")
    d.add_paragraph(summary)
    p = sec.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(); r.font.size = Pt(9)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); r._r.append(field)
    for node in list(d.styles.element.xpath(".//w:pBdr")) + list(d.element.xpath(".//w:pBdr")):
        node.getparent().remove(node)
    d.core_properties.title = title
    d.core_properties.subject = "Methods results and discussion notes for the poplar paper"
    return d
def h(d,t): d.add_heading(t,1)
def p(d,t): d.add_paragraph(t)
def table(d, headers, body, widths):
    t = d.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c,w in zip(t.columns,widths): c.width = Inches(w)
    for i,s in enumerate(headers): t.rows[0].cells[i].text=s
    for row in body:
        for c,s in zip(t.add_row().cells,row): c.text = str(s)
    for ri,row in enumerate(t.rows):
        pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit"); pr.append(cant)
        if ri == 0: pr.append(OxmlElement("w:tblHeader"))
        for ci,c in enumerate(row.cells):
            c.width= Inches(widths[ci]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr()
            borders=OxmlElement("w:tcBorders")
            for edge in ["top","left","bottom","right"]:
                b=OxmlElement("w:"+edge); b.set(qn("w:val"),"single"); b.set(qn("w:sz"),"5"); b.set(qn("w:color"),"D9D9D9"); borders.append(b)
            tcpr.append(borders)
            margins=OxmlElement("w:tcMar")
            for edge in ["top","left","bottom","right"]:
                m=OxmlElement("w:"+edge);m.set(qn("w:w"),"90");m.set(qn("w:type"),"dxa");margins.append(m)
            tcpr.append(margins)
            shade=OxmlElement("w:shd"); shade.set(qn("w:fill"),"24435A" if ri==0 else ("F2F5F7" if ri%2==0 else "FFFFFF"));tcpr.append(shade)
            for para in c.paragraphs:
                para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3)
                if ci>0: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:
                    r.font.size=Pt(9.5);r.bold=ri==0
                    r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
    d.add_paragraph().paragraph_format.space_after=Pt(0)
def sources(d, paths):
    h(d,"Files supporting this note")
    for path in paths:
        para=d.add_paragraph(path);para.paragraph_format.space_after=Pt(3)
        for r in para.runs:r.font.size=Pt(9)
def picture(d, name, caption, width=6.1):
    from publication_captions import publication_caption
    caption = publication_caption(Path(name).stem, caption)
    path=ROOT/"results"/"figures"/name
    para=d.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=para.add_run(); inline=run.add_picture(str(path),width=Inches(width))
    inline._inline.docPr.set("descr",caption)
    d.add_paragraph(caption,"Caption")
def save(d,name):
    path=OUT/name;d.save(path);print(path)

# Growth
d=make("Tree growth and carbon increment analysis",
 "This note describes growth and aboveground carbon accumulation through time in the four poplar hybrids. PAIC measures recent growth over an observed interval; MAIC measures average accumulation since planting. The current figures are descriptive, and the measurement gaps require explicit interpretation.")
h(d,"Methods")
p(d,"The trees were planted in 1992. The dataset contains 180 measured trees: Beaupre, Gibecq, Robusta and Trichobel, with 45 trees per hybrid across three experimental blocks. Measurements cover the even-numbered tree rows. The recorded series spans 1993 to 2023; there are no visits in 1999, 2000, 2003 or 2004.")
p(d,"The main preparation retains supplied tree-level carbon stocks, PAIC and MAIC. Separate audit columns reconstruct biomass using the supplied diameter allometry, convert biomass to carbon using a fraction of 0.5, and scale to area using the configured density of 166 trees per hectare. This scaling is an assumption, not a measured whole-system carbon balance.")
p(d,"PAIC is the change in aboveground carbon stock divided by elapsed years between successive visits. MAIC is carbon stock divided by years since planting. Annual plots show hybrid medians and interquartile ranges across trees. These ranges describe variability and are not confidence intervals for independent experimental replicates.")
h(d,"Main results")
annual=rows("results/tables/growth_by_year_hybrid.csv")
vals={}
for r in annual:
    if r["measurement_year"]=="2023": vals[(r["hybrid"],r["metric"])]=r["median"]
table(d,["Hybrid","PAIC","MAIC","Carbon stock"],[
 [hy,num(vals[(hy,"PAIC")]),num(vals[(hy,"MAIC")]),num(vals[(hy,"carbon_stock_mg_ha")])]
 for hy in ["Beaupre","Gibecq","Robusta","Trichobel"]],[2.0,1.4,1.4,1.7])
p(d,"The table gives 2023 medians. PAIC and MAIC are in Mg C per hectare per year; stock is in Mg C per hectare. Trichobel has the highest median values at the final visit. This descriptive ranking does not substitute for a longitudinal statistical comparison.")
d.add_page_break();h(d,"Measurement gaps and interpretation")
p(d,"The 2001 PAIC averages the 1998 to 2001 measurement interval, corresponding to growth years 1998 to 2000 under the adopted convention. Missing visits do not mechanically lower annualised PAIC: dividing by elapsed time already accounts for duration. The gap prevents identification of which individual years contributed to the low average and can smooth annual extremes. The 2005 PAIC similarly covers growth years 2002 to 2004.")
p(d,"Recalculation confirms the dip: mean PAIC across available tree records is approximately 2.65 at the 1998 visit, 0.89 at the 2001 visit and 1.77 at the 2002 visit. The supplied and stock-derived values agree closely at these visits. The cause of the slowdown remains unresolved.")
picture(d,"paic_maic_descriptive.png","Figure 1 Hybrid medians and interquartile ranges across measured trees. Lines connect visits without reconstructing unobserved annual growth.",5.5)
h(d,"Discussion and decisions still needed")
p(d,"Use PAIC to discuss changes in growth rate, and MAIC or stocks to discuss accumulated carbon. Resolve the remaining Robusta response-summary discrepancy, the treatment of missing or nonpositive observations, and allometric uncertainty before final inference. Do not interpret continued stock differences as evidence that current growth rates remain different.")
sources(d,["analysis/01_growth_descriptive.R","results/tables/growth_by_year_hybrid.csv","results/checks/tree_response_audit.csv","results/checks/response_comparison.csv"])
save(d,"01_tree_growth.docx")

# Weather
d=make("Weather relationships with tree growth",
 "This note records the current hybrid-specific weather models for PAIC and the limitations that matter for paper drafting. The baseline models identify statistical associations with multivariate weather gradients; they do not yet explain the low growth recorded at the 2001 visit.")
h(d,"Methods")
p(d,"The baseline analysis uses the legacy hybrid PAIC summaries for continuity with the earlier weather analysis. These have been compared with tree-level hybrid medians. Beaupre, Gibecq and Trichobel agree closely; a remaining Robusta discrepancy needs resolution. Only positive, observed PAIC values enter the Gamma models.")
p(d,"A measurement in year t is paired with weather for growth year t minus 1. Seasons are winter December to February, spring March to May, summer June to August and autumn September to November; half-years are October to March and April to September. Winter windows cross the calendar-year boundary.")
p(d,"Temperature, soil temperature and wind speed are averaged; rainfall, sunshine, frost-day counts, growing-day counts, growing degree days and windy-day counts are summed within windows. Variable coverage is exported. Partial variable records currently use available days; a final coverage policy has not been fixed.")
p(d,"Six separate principal component analyses describe the seasonal or half-year weather. Each is shared across hybrids, uses standardised weather variables on its complete-case basis, and retains enough axes to explain at least 80 percent of weather variance. For each hybrid and period, PAIC is modelled with a Gamma distribution and log link using retained component scores.")
h(d,"Current model results")
wm=rows("results/tables/baseline_weather_model_summary.csv")
table(d,["Hybrid","Summer observations","Summer PCs","Deviance reduction"],[
 [r["hybrid"],r["n"],r["n_pcs"],f'{float(r["deviance_reduction"])*100:.1f}%']
 for r in wm if r["weather_period"]=="summer"],[1.7,1.6,1.1,2.1])
p(d,"All exported baseline fits converged. Deviance reduction is an in-sample measure, not predictive validation or the fraction of growth causally explained by weather. Models fitted to different observation sets must not be ranked directly by AIC.")
d.add_page_break();h(d,"The unresolved slowdown around the 2001 visit")
p(d,"PAIC at the 2001 visit averages growth over 1998 to 2000, but the current weather join supplies weather for 2000 only. This mismatch matters when interpreting that point. No interval-aware weather analysis has yet been completed, and the existing model should not be used to attribute the three-year average to weather in one year.")
ww=rows("data/processed/weather_wide.csv")
table(d,["Growth year","Summer half rainfall mm","Summer half sunshine hours"],[
 [r["growth_year"],num(r["summerhalf_rainfall"],1),num(r["summerhalf_sunshine"],1)]
 for r in ww if r["growth_year"] in ["1997","1998","1999","2000","2001"]],[1.2,2.5,2.8])
p(d,"The local records show relatively low summer-half sunshine in 1998 and higher summer-half rainfall in 2000. These are observations for investigation, not established explanations of the growth decline. They do not resolve which years within the measurement gap were most influential.")
h(d,"Discussion points")
p(d,"Interpret each component using its loadings, rather than assigning a physical meaning from its sign or number. Component signs are arbitrary, and a one-unit score is not a one-degree temperature or one-millimetre rainfall change. Any quoted association should name the hybrid, weather period, component interpretation and uncertainty.")
p(d,"The treatment analysis absorbs shared temporal conditions with categorical interval effects. Its interval coefficient cannot distinguish weather from stand development, management changes or other shared causes. Separate weather modelling can investigate associations, while treatment by weather interactions would be a distinct, more demanding question.")
h(d,"Next analysis decisions")
p(d,"Compare annual-interval models with an interval-aware sensitivity analysis. A candidate approach is to predict annual growth using annual weather and then average predictions over each observed growth interval; this is not generally identical to predicting from averaged weather under a nonlinear link. Assess age or time trends, residual dependence, influential observations and fair temporal validation before describing forecasts as reliable.")
sources(d,["analysis/02_weather_paic.R","scripts/prep/02_weather.R","scripts/prep/04_join_and_pca.R","results/tables/baseline_weather_model_summary.csv","results/tables/baseline_weather_model_coefficients.csv","data/processed/weather_pca_loadings.csv","results/checks/weather_coverage.csv"])
save(d,"02_weather_and_PAIC.docx")

# Treatment
d=make("Cropping treatment effects and growth legacy",
 "This note records the PAIC comparison among continuously cropped, alternately cropped and continuously fallow alleys, during cultivation and after cessation. Continuous cropping reduced PAIC during active cultivation. After cessation, a temporary growth advantage over former continuous fallow is suggestive of release from crop competition, with the strongest interval-specific evidence in growth years 2008 and 2009.")
h(d,"Experimental design and response")
p(d,"Within each of three blocks, six alleys were allocated in adjacent pairs to the three cropping treatments. A measurement row lay between each paired set of alleys. Continuous fallow had both alleys fallow; continuous cropping had both cropped; alternate cropping had one cropped and one fallow, reversed each year. Alternate cropping is not a second wholly fallow treatment.")
p(d,"All available measured trees are retained. There are nine treatment measurement rows with five trees of each of four hybrids per row, giving 180 tree IDs. The design describes guard trees at row ends. Each row by hybrid by interval mean is equally weighted in the model; the five neighbouring trees are subsamples, not independent treatment replicates.")
h(d,"Growth intervals and cessation")
p(d,"Cropping stopped after the 2004 harvest. The primary active model includes observed intervals through the 2002 visit. The primary post-cropping model begins with the 2005 to 2006 interval, representing growth year 2005. The 2002 to 2005 boundary interval includes growth years 2002 to 2004 under the adopted timing convention. It is excluded from both primary fits and included with active cropping in a sensitivity analysis.")
p(d,"The initial visit has no preceding increment and is excluded. The low PAIC at the 2001 visit is a three-year average across growth years 1998 to 2000. Missing visits limit temporal resolution; they do not themselves cause low annualised growth.")
h(d,"Statistical model")
p(d,"Separate Gaussian mixed models include fixed block and the full treatment by hybrid by interval interaction. Random intercepts represent treatment row, row by hybrid and row by interval. Block remains a design adjustment irrespective of significance. Treatment comparisons average equally across hybrids and observed intervals; this differs from duration-weighting the intervals.")
p(d,"Overall treatment contrasts use Bonferroni adjustment for three pairs. Interval-specific inference uses Holm adjustment across all three pairs and all intervals within each phase: 21 active, 54 post-cropping and 24 boundary-sensitivity tests. Confidence intervals in the marked figures are pointwise; markers use the phase-wide Holm tests. The three fits are not pooled into a single multiplicity family because the primary phases address separate questions and the third fit is a sensitivity analysis.")
d.add_page_break();h(d,"Main results")
body=[]
for phase,display in [("active","Active"),("post","Post")]:
    for r in rows(f"results/tables/treatment_PAIC_{phase}_overall_contrasts.csv"):
        body.append([display,r["contrast"],num(r["estimate"],3),f'{float(r["p.value"]):.4f}'])
table(d,["Phase","Treatment difference","PAIC difference","Adjusted p"],body,[0.65,3.4,1.25,1.2])
p(d,"PAIC differences are Mg C per hectare per year. The active treatment omnibus p-value is 0.0057. The post-cropping treatment omnibus p-value is 0.0806; the average post-cropping pairwise contrasts are uncertain and do not exclude shorter-lived differences.")
p(d,"Including the boundary interval gives continuous cropping minus continuous fallow of -0.445, adjusted p = 0.0075. The corresponding continuous cropping minus alternate cropping contrast is -0.260, adjusted p = 0.0506. The continuous-cropping versus continuous-fallow conclusion is consistent; the other active average contrast is sensitive near the conventional threshold.")
h(d,"Differences among hybrids")
p(d,"Hybrid main effects are significant in both primary models, with p below 0.001. Treatment by hybrid interactions are nonsignificant: active p = 0.613 and post p = 0.570. The three-way interactions with interval are also nonsignificant. We therefore display treatment comparisons averaged across hybrids, without claiming that hybrids have equal PAIC or identical treatment responses.")
h(d,"Post cropping intervals surviving adjustment")
inf=rows("results/tables/treatment_PAIC_post_interval_inference.csv")
table(d,["Growth year","Comparison","Difference","Holm p"],[
 [r["growth_year"],r["contrast"],num(r["estimate"],3),f'{float(r["p_holm_phase"]):.4f}']
 for r in inf if float(r["p_holm_phase"])<.05],[0.9,3.5,1.1,1.0])
p(d,"These correspond to the 2009 and 2010 visits. The broader positive pattern over growth years 2007 to 2010 is descriptive, not significant in every year after phase-wide correction. The growth-year-2013 difference does not survive correction; no interval has been removed as a presumed fluke.")
sources(d,["analysis/03_treatment_active.R","scripts/treatment_presentation.R","analysis/03b_refresh_treatment_figures.R","results/tables/treatment_PAIC_*_overall_contrasts.csv","results/tables/treatment_PAIC_*_interval_inference.csv"])
d.add_page_break();h(d,"Discussion of the post cropping legacy")
picture(d,"treatment_PAIC_post_means.png","Figure 1 Model-adjusted PAIC averaged across hybrids after cessation. Asterisks mark continuous cropping versus continuous fallow differences after phase-wide Holm correction. Error bars are pointwise 95 percent confidence intervals.",6.1)
p(d,"Suggested drafting language: Following cessation of cropping, trees formerly under continuous cropping showed a temporary PAIC advantage over trees formerly under continuous fallow, with the strongest statistical evidence in growth years 2008 and 2009. This pattern is consistent with release from crop competition, but does not establish that mechanism.")
p(d,"Post-cropping contrasts include consequences of earlier differences in size and condition. Weather, stand development and other temporal influences may contribute to the trajectories. A nonsignificant contrast does not date disappearance of a legacy or demonstrate equivalence between treatments.")
h(d,"Active cropping convergence and remaining decisions")
p(d,"During the 1998 to 2001 interval, PAIC fell in all treatments and the treatment differences narrowed. Continuous fallow minus continuous cropping declined from about 0.83 at the 1998 visit to 0.20 at the 2001 visit, then rose to 0.29 at the 2002 visit. The weather explanation remains untested with interval-aligned predictors.")
p(d,"All current fits converged without singularity. Residual plots show some unequal spread and tail departures; variance assumptions and remaining serial dependence require review before final inferential wording. Three blocks limit independent treatment replication. A biological equivalence threshold would be needed to estimate when a legacy becomes negligible.")
save(d,"03_cropping_treatment_and_legacy.docx")

# Soil: use the current depth-specific figure and 0-30 cm supplementary tables.
import runpy
runpy.run_path(str(ROOT / "scripts" / "update_soc_note.py"), run_name="__main__")

# Scenarios
d=make("Carbon scenario analysis planning",
 "This note records the scenario component planned for the paper and the decisions needed before calculation. No new projections have been agreed or calculated. The intended deliverable is a transparent Excel workbook that colleagues can use without R.")
h(d,"Purpose and available evidence")
p(d,"Explore how measured tree carbon accumulation and soil-carbon information might inform explicit land-use or management scenarios. The annual stock, PAIC and MAIC summaries and complete-profile soil summaries are available as inputs. Historical scenario sheets remain as context; they are not inputs to the current preparation pipeline.")
h(d,"Methods to agree before building")
p(d,"First specify geography, eligible land categories, the area-data source and reference year. Define whether the hectare refers to the whole system or the tree component, and distinguish eligible land from adopted land. Specify adoption through time, projection horizon, hybrid composition, management and stand-age cohorts.")
p(d,"Choose how PAIC, MAIC and stock observations support each calculation. PAIC is an interval growth rate, MAIC is cumulative stock divided by age, and stock is an accumulated amount. They cannot be interchanged as annual rates. A constant-rate extrapolation would need explicit justification and limits.")
p(d,"Keep tree and soil contributions separate until their area and time conventions match. Any soil contribution requires a justified counterfactual and accumulation timescale. A cross-sectional stock difference must not automatically be divided by stand age and presented as a measured sequestration rate.")
h(d,"Current results and discussion")
p(d,"There are no new scenario results to report. Potential interpretation should be framed as conditional outcomes under stated assumptions, not a forecast of achieved removal. Avoid double counting stock and annual increments, tree-row and alley areas, or annual and cumulative totals.")
h(d,"Workbook requirements")
p(d,"Use clearly identified editable inputs, units beside each input, transparent formulas, source references and separate annual and cumulative outputs. If converting carbon to carbon dioxide, state the molecular-weight conversion explicitly. Include uncertainty or sensitivity analysis for the assumptions that materially affect conclusions.")
p(d,"Agree scenario labels, baseline or counterfactual, eligible-area fractions, adoption schedules, hybrid-specific growth trajectories, soil treatment, and the handling of harvesting, mortality and other carbon pools before producing paper figures. Avoid adding unsupported precision to uncertain inputs.")
sources(d,["scenarios/README.md","data/processed/annual_hybrid.csv","data/processed/annual_hybrid_treatment.csv","data/processed/soil_stock_summary.csv","data/raw/soil_and_legacy_scenarios.xlsx"])
save(d,"05_carbon_scenario_planning.docx")


