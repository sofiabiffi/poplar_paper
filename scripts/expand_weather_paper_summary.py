from pathlib import Path
from datetime import datetime
import csv,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
d=Document(path);title='Weather associations with annual carbon increment'
anchor=next(p for p in d.paragraphs if p.text==title)
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'{path.stem}_before_expanded_absolute_summary_{datetime.now():%Y%m%d_%H%M%S}.docx')
# Replace the authorised final summary only; retain preceding notes verbatim.
body=d._element.body;node=anchor._p
while node is not None and node.tag!=qn('w:sectPr'):
    nxt=node.getnext();body.remove(node);node=nxt
start=len(d.paragraphs)
def p(text):return d.add_paragraph(text)
def h(text):return d.add_heading(text,2)
d.add_heading(title,1)
p('Paper analysis summary | Updated 21 October 2026. This self-contained account focuses on PAIC and weather associations in absolute carbon units. Percentage and standardised results remain in the earlier working record, but are not the proposed paper focus.')
h('Purpose and response data')
p('We ask whether observed fluctuations in carbon increment are associated with summer weather, and whether these associations differ among Beaupre, Gibecq, Robusta and Trichobel. This is a supporting explanatory analysis, rather than a forecasting exercise or a test that climate change caused the growth differences.')
p('PAIC is the carbon increment between successive measurements divided by the elapsed years, expressed in Mg C ha-1 yr-1. The model uses positive tree-level median PAIC for each hybrid in 26 common measurement years: 104 hybrid summaries. The four hybrids share weather in a given year, so these are 26 distinct weather exposures, not 104 independent weather replicates. The final analysis uses tree medians consistently, rather than the legacy response series used in some earlier models.')
h('Previous year weather and sampling gaps')
p('A measurement in year t is matched to weather from growth year t minus 1. The retained predictors are June-August mean air temperature and total rainfall; for example, 2002 PAIC is matched to summer 2001. PAIC following a multi-year gap is an average annual increment over that interval, but its assigned weather still describes only the final growth year. Annualisation accounts for interval duration; it cannot recover the timing of growth in the missing years. We retain this convention and disclose the limitation.')
h('How weather variables were selected')
p('We first formulated biological hypotheses about warmth, water supply, light availability and stress. Nine candidate summaries represented them: spring and summer air temperature, rainfall and sunshine (six variables); spring air-frost and ground-frost day counts (two); and the longest April-September dry spell (one). Spring was March-May and summer June-August. Frost days had minimum air or grass temperature below 0 degrees C; dry days had rainfall below 1 mm.')
p('For this exploratory screening only, a separate quadratic age trend was removed from log PAIC for each hybrid. Departures were divided by each hybrid residual standard deviation and then averaged with equal hybrid weights. We examined Pearson correlations, which describe linear associations, and Spearman correlations, which compare ranks. We also checked the effect of omitting individual years, using linear detrending instead, and examining hybrids separately.')
p('The common screen required complete daily coverage for all nine measures, leaving 24 years. Summer rainfall and temperature had the clearest positive Pearson correlations with the shared departures, r = 0.30 and 0.37. Rainfall was more consistent across rank and influence checks; temperature was more sensitive to the earliest observation. The other screened measures had weaker associations. Both summer predictors were retained, with no further variable search. The final model has 26 complete years because it needs only these two weather measures. Selection used the same study data, so subsequent associations remain exploratory.')
d.add_page_break();d.add_heading('Statistical model and adjusted comparisons',1)
h('Model formulation')
p('The final model is fitted directly to original positive PAIC, without standardising it or removing an age trend first. Each hybrid has its own baseline, curved age trajectory, temperature association and rainfall association. The mean model is:')
eq=p('log(μ) = α + β₁a + β₂a² + γT* + δP*')
for r in eq.runs:r.font.name='Cambria Math';r.font.size=Pt(14)
p('Here μ is the expected PAIC for a particular hybrid and measurement year. All five coefficients are specific to that hybrid. The age variable is a = (measurement age - 15)/10. Temperature is T* = previous-summer mean temperature - 16 degrees C. Rainfall is P* = (previous-summer rainfall - 200 mm)/100 mm. These fixed centres improve numerical stability; they do not alter the fitted PAIC or weather contrasts.')
p('The response is modelled with a Gamma distribution and a log link. This allows positive PAIC with unequal scatter and guarantees positive fitted mean values. The model describes the logarithm of the mean PAIC; it does not fit a linear model to log-transformed observations. Each hybrid is fitted separately, equivalently allowing hybrid interactions with every age and weather term in a joint mean model. Paired resampling, described below, preserves the shared years for uncertainty and hybrid comparisons.')
h('What absolute units do and do not mean')
p('Absolute means the original carbon units, not that growth must increase or that negative differences are removed. PAIC can fall from, for example, 2 to 1 Mg C ha-1 yr-1 while remaining positive. The age terms can describe a declining or curved trajectory, and the actual weather sequence can add rises and falls. The present fitted temperature and rainfall associations are positive, but that does not force fitted PAIC to increase through time. Gamma models cannot represent zero or negative PAIC; only positive observations enter this analysis.')
p('A weather-effect graph is not a time-trend graph. It compares two weather values while holding other conditions comparable. A fitted time series instead uses each year\'s actual age and weather. Consequently, the positive weather contrasts below are not estimates of the net PAIC change across the study period and do not imply that declining growth has been reversed by warming.')
h('How absolute weather differences are calculated')
p('For temperature, compare 15.31 with 16.31 degrees C, holding rainfall at its observed mean. For rainfall, compare 122.4 with 222.4 mm, holding temperature at its mean. For each hybrid, calculate mean PAIC at both settings over the same 26 observed ages, then subtract the lower-weather mean from the higher-weather mean. The resulting differences are in Mg C ha-1 yr-1. They are specific to these reference settings and age distribution; because the link is nonlinear, there is no single constant absolute effect at every age and weather level.')
h('Bootstrap explanation and limitations')
p('We constructed 1,000 resampled versions of the record by drawing blocks of three neighbouring visits, with all four hybrids kept together. Some visits were repeated and others omitted. We refitted the complete model each time, recalculated weather differences and compared hybrids within the same resample. The middle 95% of estimates forms the pointwise uncertainty interval. This is an approximate uncertainty check, not additional independent data.')
p('Blocks retain some short temporal dependence but follow observed visits across sampling gaps; the circular procedure also joins the record end to its start. Three visits is a provisional block length. Some resamples poorly represent early ages and produce extreme adjusted estimates when averaged over the full age range. No samples were discarded, but percentile intervals do not resolve that instability. The bootstrap does not account for the earlier weather-variable screening.')
d.add_page_break();d.add_heading('Absolute results and interpretation',1)
h('Adjusted differences in PAIC')
p('Entries are estimated differences in Mg C ha-1 yr-1, with pointwise 95% block-bootstrap intervals in brackets. A positive value means higher fitted PAIC at the higher weather setting, after the specified age and other-weather adjustment.')
with (root/'results/tables/absolute_weather_adjusted_differences.csv').open(newline='') as f:rows=list(csv.DictReader(f))
t=d.add_table(rows=1,cols=3);t.autofit=False;widths=[1.05,2.72,2.72]
for i,txt in enumerate(['Hybrid','Temperature difference of 1 degree C','Rainfall difference of 100 mm']):t.rows[0].cells[i].text=txt
for name in ['Beaupre','Gibecq','Robusta','Trichobel']:
    vals=[]
    for v in ['temperature','rainfall']:
        r=next(x for x in rows if x['hybrid']==name and x['term']==v)
        vals.append(f"{float(r['absolute_difference']):+.2f} [{float(r['lower']):+.2f}, {float(r['upper']):+.2f}]")
    for c,txt in zip(t.add_row().cells,[name]+vals):c.text=txt
for ri,row in enumerate(t.rows):
    for ci,c in enumerate(row.cells):
        c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        pr.append(borders)
        if ri==0:
            s=OxmlElement('w:shd');s.set(qn('w:fill'),'E8EDF0');pr.append(s)
        for para in c.paragraphs:
            para.paragraph_format.space_before=Pt(5);para.paragraph_format.space_after=Pt(5)
            for r in para.runs:r.font.size=Pt(10);r.bold=ri==0
p('Trichobel has the largest absolute point estimate for both weather contrasts. This is compatible with more additional carbon increment under warmer or wetter summers within the observed range, but higher baseline PAIC contributes to the size of an absolute difference. All 12 direct paired hybrid-difference intervals include zero. Thus the analysis does not clearly establish that Trichobel has a larger weather response than the other hybrids, and a nonsignificant difference is not evidence that responses are equal.')
p('Using a linear rather than quadratic age trajectory gives temperature differences of approximately +0.45, +0.35, +0.48 and +0.52 for Beaupre, Gibecq, Robusta and Trichobel, respectively. This sensitivity preserves the broad absolute ranking, but is not independent confirmation of the model or mechanism.')
h('Discussion for the paper')
p('Summer warmth and rainfall are positively associated with adjusted PAIC in this record and may contribute to the observed growth fluctuations. Trichobel maintained higher growth over much of the study and has the largest estimated absolute weather differences. This is consistent with the biological hypothesis that it copes well with changing conditions. However, the model does not demonstrate that climate change caused its study-period advantage or that it is more tolerant of heat, drought or flooding. A positive association with a seasonal average is not a direct test of resistance to extreme events.')
p('All trees belong to one planting cohort, so age and calendar time move together while weather and management history also change. Age adjustment can absorb variation associated with gradual warming or hybrid divergence, and the data cannot cleanly allocate that variation among causes. Residual temporal patterns, limited independent weather years, sampling gaps and unstable bootstrap fits add uncertainty. These results should therefore support a restrained discussion of possible explanations, rather than carry the main explanation for differences among hybrids.')
h('Scope and reproducibility')
p('For the paper, prioritise original PAIC curves and absolute differences with uncertainty. Keep percentage and standardised analyses as working background, not additional headline results. The present extension changes the interpretation and reporting scale, not the selected summer predictors. No new model was fitted for this explanatory revision.')
p('Analysis: analysis/13_absolute_weather_paic.R. Results: results/tables/absolute_weather_adjusted_differences.csv, absolute_weather_paired_differences.csv, absolute_weather_linear_age_sensitivity.csv and absolute_weather_bootstrap_tail_audit.csv. Figures: absolute_weather_curves and absolute_weather_differences, PNG and PDF. Earlier pages retain the exploratory history; this section is the current account for paper drafting.')
for para in d.paragraphs[start:]:
    if para.style.name not in ['Heading 1','Heading 2'] and para is not eq:
        para.paragraph_format.space_after=Pt(7);para.paragraph_format.line_spacing=1.05
        for r in para.runs:
            if r.font.name!='Cambria Math':r.font.size=Pt(10.5)
    for r in para.runs:r.font.color.rgb=RGBColor(0,0,0)
d.save(path);print(path)
