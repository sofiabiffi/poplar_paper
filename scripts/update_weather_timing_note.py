"""Expand weather timing and explain the age curves, without changing estimates."""
from pathlib import Path
from datetime import datetime,date,timedelta
import calendar,csv,math,shutil,sys
from docx import Document
from docx.shared import Inches,Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(sys.argv[1]) if len(sys.argv)>1 else Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
def read(rel):
    with (root/rel).open(encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))
windows=read('data/processed/weather_windows.csv')
expected={
 'spring':('2001-03-01','2001-05-31'), 'summer':('2001-06-01','2001-08-31'),
 'autumn':('2001-09-01','2001-11-30'), 'winter':('2000-12-01','2001-02-28'),
 'winterhalf':('2000-10-01','2001-03-31'), 'summerhalf':('2001-04-01','2001-09-30')}
audit=[]
for period,(start,end) in expected.items():
    w=next(w for w in windows if w['measurement_year']=='2002' and w['weather_period']==period)
    assert w['growth_year']=='2001' and w['start_date']==start and w['end_date']==end
    audit.append(dict(measurement_year=2002,growth_year=2001,window=period,start_date=start,end_date=end,n_days=(date.fromisoformat(end)-date.fromisoformat(start)).days+1))
daily=read('data/processed/weather_daily.csv')
summer=[r for r in daily if '2001-06-01'<=r['date']<='2001-08-31']
w=next(w for w in windows if w['measurement_year']=='2002' and w['weather_period']=='summer')
assert len(summer)==92
assert abs(sum(float(r['rainfall']) for r in summer)-float(w['rainfall']))<1e-8
assert abs(sum(float(r['temp']) for r in summer)/92-float(w['temp']))<1e-8
trees=read('data/processed/tree_growth.csv')
gap=[r for r in trees if r['measurement_year']=='2001' and r['PAIC']]
assert gap and all(float(r['previous_measurement_year'])==1998 and float(r['interval_years'])==3 and r['growth_year']=='2000' for r in gap)
assert all(r['growth_year']=='2022' for r in trees if r['measurement_year']=='2023')
with (root/'results/checks/weather_timing_methods_audit.csv').open('w',encoding='utf-8',newline='') as f:
    writer=csv.DictWriter(f,fieldnames=list(audit[0]));writer.writeheader();writer.writerows(audit)

d=Document(path)
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'02_weather_and_PAIC_before_timing_explanation_{datetime.now():%Y%m%d_%H%M%S}.docx')
def find(prefix):
    found=[p for p in d.paragraphs if p.text.startswith(prefix)]
    assert len(found)==1,(prefix,len(found))
    return found[0]
def put_before(anchor,text,style=None):
    return anchor.insert_paragraph_before(text,style)
def heading_before(anchor,text,page=False):
    p=put_before(anchor,text,'Heading 1');p.paragraph_format.page_break_before=page;return p
def remove(p):p._p.getparent().remove(p._p)
anchor=find('Comparing the weather hypotheses')
# Allow deliberate reruns without accumulating explanation pages.
old=[p for p in d.paragraphs if p.text=='Timing of PAIC and weather windows']
if old:
    body=d.element.body;start=list(body).index(old[0]._p);end=list(body).index(anchor._p)
    for node in list(body)[start:end]:body.remove(node)
heading_before(anchor,'Timing of PAIC and weather windows',True)
put_before(anchor,'The year labels identify measurements, while the weather must describe the growth that those measurements capture. Let t be the recorded measurement year and s the previous observed measurement year. Under the adopted dormant-season measurement convention, the most recent completed growth year is g = t − 1. This is an alignment of growth and weather, not evidence of a delayed one-year physiological response to weather.')
put_before(anchor,'PAIC recorded at t represents the annualised carbon-stock increment over the interval from s to t: PAICₜ = (Cₜ − Cₛ)/(t − s). For consecutive annual visits, s = t − 1. After a measurement gap, s is the last year actually observed. The weather models use supplied annualised tree-level PAIC values, summarised by hybrid medians; separately recalculated stock differences are preparation audits rather than replacements of the supplied response.')
put_before(anchor,'The original online discussion used a March 2002 visit as the example: the increment since the preceding annual visit mainly represents the 2001 growing season. Relevant summer weather is therefore June–August 2001. The winter preceding that growing season is December 2000–February 2001. December 2001–February 2002 is the winter after that growth season, immediately before the March 2002 visit; it is not the “previous winter” window defined in these scripts.')
headers=['Weather window','Dates relative to measurement year t','Example for t = 2002']
rows=[
 ['Spring','1 March–31 May of t − 1','1 March–31 May 2001'],
 ['Summer','1 June–31 August of t − 1','1 June–31 August 2001'],
 ['Autumn','1 September–30 November of t − 1','1 September–30 November 2001'],
 ['Previous winter','1 December of t − 2 to the last day of February of t − 1','1 December 2000–28 February 2001'],
 ['Winter half-year','1 October of t − 2 to 31 March of t − 1','1 October 2000–31 March 2001'],
 ['Summer half-year and dry-spell window','1 April–30 September of t − 1','1 April–30 September 2001'],
]
t=d.add_table(rows=1,cols=3);t.autofit=False
widths=[1.35,2.65,2.6]
for c,w in zip(t.columns,widths):c.width=Inches(w)
for c,v in zip(t.rows[0].cells,headers):c.text=v
for values in rows:
    for c,v in zip(t.add_row().cells,values):c.text=v
for ri,row in enumerate(t.rows):
    pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
    if ri==0:pr.append(OxmlElement('w:tblHeader'))
    for ci,c in enumerate(row.cells):
        c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        tcpr.append(borders);margins=OxmlElement('w:tcMar')
        for edge in ['top','left','bottom','right']:
            e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');margins.append(e)
        tcpr.append(margins)
        if ri==0:
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7EDF2');tcpr.append(sh)
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            for r in p.runs:r.font.size=Pt(9.5);r.bold=(ri==0)
anchor._p.addprevious(t._tbl)
window_caption=put_before(anchor,'All boundaries are inclusive. Winter length follows the actual calendar, including 29 February in leap growth years. A calendar-year winter made from January, February and December of the same year would mix winters on either side of the growth season; the code instead constructs continuous date ranges from daily records.','Caption')
window_caption.paragraph_format.space_before=Pt(6)
put_before(anchor,'The current eight hypotheses use spring rainfall, temperature, sunshine and frost; summer rainfall, temperature and sunshine; and the April–September dry spell. Autumn, previous winter and winter-half-year variables are prepared for context but are not tested in the current eight-model comparison. April–September defines the dry-spell search interval; it does not replace June–August as the summer window. A combined previous-winter-plus-summer model and a calendar-year annual model are not part of the present analysis.')

heading_before(anchor,'Daily aggregation and Methods wording',True)
put_before(anchor,'Rainfall (mm) and sunshine (hours) are daily sums within each window. Temperature is the mean of daily mean-air-temperature values, not an unweighted average of monthly means; summer has 92 days. Spring frost counts use daily minimum air or grass temperature below 0 °C. The dry spell is the longest run of days with rainfall <1 mm within April–September.')
put_before(anchor,'Every expected date and required daily value must be available. Completeness is checked on the prepared daily series; these analyses introduce no further gap-filling. The eight hypotheses use 24 identical complete-weather years. The reporting summer model uses all 26 years with complete June–August rainfall and temperature.')
put_before(anchor,'Measurement gaps do not change the offset. PAIC measured in 2001 averages the stock change since 1998 over three years, representing growth years 1998–2000. Weather still describes 2000, including June–August 2000 for summer; it is not averaged across the three growth years. Annualisation cannot recover growth within missing years. The final 2023 observation similarly uses summer 2022, not summer 2023.')
put_before(anchor,'Suggested Methods paragraph','Heading 2')
put_before(anchor,'Periodic annual carbon increment (PAIC) was assigned to the year of the later DBH measurement and annualised over the number of years since the preceding available measurement. Weather predictors were aligned with the preceding growth year (g = t − 1, where t denotes measurement year), consistent with dormant-season measurements recording growth completed in the preceding growing season. Spring predictors were summarised from 1 March to 31 May and summer predictors from 1 June to 31 August of g. Seasonal rainfall and sunshine were summed from daily records, and air temperature was averaged over daily mean values. Spring frost counts used daily minimum air or grass temperature below 0 °C; longest dry spell was calculated over 1 April–30 September of g using a daily rainfall threshold of <1 mm. Only complete daily windows were used. For multi-year measurement intervals, PAIC remained annualised over the full interval but was paired with weather from the final growth year, because the timing of growth within unmeasured years was unknown.')
put_before(anchor,'Timing qualification and source trail','Heading 2')
put_before(anchor,'The earlier co-author comment identifies March measurements and assumes earlier observers followed that practice; the thesis describes November–March. The tree table contains year labels, not individual visit dates. The code applies t − 1 consistently rather than correcting individual November or March visits. Describe this adopted convention without claiming that every historical visit date was individually verified.')
put_before(anchor,'Verified sources: scripts/config.R; scripts/prep/01_tree_growth.R and 02_weather.R; analysis/10_weather_departure_screen.R, 14_raw_paic_weather_final.R and 17_weather_hypothesis_models.R. The 2002 windows and daily summer aggregation were checked directly; dates are recorded in results/checks/weather_timing_methods_audit.csv. The earlier discussion and co-author example are in “Jyoti Code Walkthrough Thesis Analysis” (https://chatgpt.com/c/6a1ebf0c-89a8-8391-abab-ad1e97d666a1).','Caption')

p=find('Spring means March–May and summer means June–August.')
p.text=p.text.replace('No missing daily weather is imputed.','No further daily weather imputation is introduced in these analyses; completeness is assessed on the prepared series.')
matches=[p for p in d.paragraphs if p.text.startswith('The fitted age trajectory declines in Beaupre') or p.text.startswith('The curves show the smooth background PAIC pattern')]
assert len(matches)==1
p=matches[0]
for prefix in ['Beaupre shows a gradual decline.','PAIC is annual carbon addition:']:
    for oldp in [p for p in d.paragraphs if p.text.startswith(prefix)]:remove(oldp)
p.text='The curves show the smooth background PAIC pattern after accounting for summer weather. They are drawn by varying age across its observed range while holding rainfall and temperature at their averages, so individual wet, dry, warm or cool years do not make the curves fluctuate.'
p2=d.add_paragraph('Beaupre shows a gradual decline. Gibecq is fairly flat initially, then declines. Robusta rises to a broad peak before declining. Trichobel rises more strongly to a later peak, then declines. These are fitted background patterns during this study.')
p._p.addnext(p2._p)
p3=d.add_paragraph('PAIC is annual carbon addition: declining PAIC means less carbon added each year, not necessarily carbon loss. Age also tracks calendar time, competition and management; the curves do not isolate biological ageing. This plot is optional working material for interpreting the model.')
p2._p.addnext(p3._p)
# Move implementation pointers to the dedicated Methods pages above.
for oldp in [p for p in d.paragraphs if p.text.startswith('Reproducibility:')]:remove(oldp)
p=find('Use Figure 1 to show the hypothesis comparison')
p.text=p.text.replace('Figure S1 shows the fitted quadratic age shapes at fixed weather.','Figure S1 explains the fitted background age patterns at fixed weather; it is optional supporting material rather than a required main-paper result.')
d.save(path)
print('Updated the date-window Methods explanation and age-curve interpretation; model estimates unchanged.')
