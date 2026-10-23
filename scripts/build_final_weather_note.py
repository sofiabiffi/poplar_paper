"""Build the current weather note only, preserving the previous Word file first.

Run from the workspace with the bundled Python after analysis/14_raw_paic_weather_final.R
and its companion checks. The original build_analysis_notes.py is historical for this note.
"""
from pathlib import Path
from datetime import datetime
import csv
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'manuscript/analysis_notes/02_weather_and_PAIC.docx'
history = OUT.parent / '_history'
history.mkdir(exist_ok=True)
backup = history / f'02_weather_and_PAIC_before_raw_revision_{datetime.now():%Y%m%d_%H%M%S}.docx'
shutil.copy2(OUT, backup)
d = Document(OUT)
for e in list(d.element.body):
    if e.tag != qn('w:sectPr'):
        d.element.body.remove(e)
sec = d.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.75)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
    style = d.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0,0,0)
d.styles['Normal'].font.size = Pt(11)
d.styles['Normal'].paragraph_format.space_after = Pt(7)
d.styles['Normal'].paragraph_format.line_spacing = 1.06
d.styles['Title'].font.size = Pt(21)
d.styles['Heading 1'].font.size = Pt(13)
d.styles['Heading 1'].paragraph_format.space_before = Pt(10)
d.styles['Heading 1'].paragraph_format.space_after = Pt(6)
d.styles['Caption'].font.size = Pt(9)
d.styles['Caption'].font.bold = False
d.styles['Caption'].paragraph_format.space_after = Pt(7)
for node in list(d.styles.element.xpath('.//w:pBdr')) + list(d.element.xpath('.//w:pBdr')):
    node.getparent().remove(node)

def rows(stem):
    with (ROOT/f'results/tables/final_weather_{stem}.csv').open(encoding='utf-8-sig',newline='') as f:
        return list(csv.DictReader(f))
def p(text,style=None): return d.add_paragraph(text,style)
def h(text): d.add_heading(text,1)
def page(title):
    heading=d.add_heading(title,1)
    heading.paragraph_format.page_break_before=True
def fig(stem,caption,width=6.6):
    from publication_captions import publication_caption
    caption = publication_caption('final_weather_'+stem, caption)
    para=p(''); para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    inline=para.add_run().add_picture(str(ROOT/f'results/figures/final_weather_{stem}.png'),width=Inches(width))
    inline._inline.docPr.set('descr',caption)
    p(caption,'Caption')
def table(headers,body,widths):
    t=d.add_table(rows=1,cols=len(headers)); t.autofit=False
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for c,w in zip(t.columns,widths): c.width=Inches(w)
    for c,s in zip(t.rows[0].cells,headers): c.text=s
    for values in body:
        for c,s in zip(t.add_row().cells,values): c.text=str(s)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0: trpr.append(OxmlElement('w:tblHeader'))
        for ci,c in enumerate(row.cells):
            c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                b=OxmlElement('w:'+edge); b.set(qn('w:val'),'single');b.set(qn('w:sz'),'5');b.set(qn('w:color'),'D9D9D9');borders.append(b)
            pr.append(borders)
            fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7EDF2' if ri==0 else 'FFFFFF');pr.append(fill)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            pr.append(mar)
            for para in c.paragraphs:
                para.paragraph_format.space_before=Pt(2);para.paragraph_format.space_after=Pt(2)
                if ci: para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                for r in para.runs:r.font.size=Pt(9.5);r.bold=(ri==0)
    p('').paragraph_format.space_after=Pt(0)

p('Weather associations with observed PAIC','Title')
p('Current analysis for paper drafting   |   21 October 2026','Subtitle')
p('Summer rainfall is positively associated with PAIC in all four hybrids after allowing a curved age trajectory associated with stand development. Trichobel has the largest estimated absolute rainfall and temperature associations, but the uncertainty does not establish larger responses than in the other hybrids. These exploratory results suggest possible weather contributions to growth, not a causal explanation of hybrid performance.')
h('Methods and data alignment')
p('The response is original PAIC, in Mg C ha⁻¹ yr⁻¹: the median of the available tree values within each hybrid and measurement year. Beaupre, Gibecq, Robusta and Trichobel are analysed separately. These summaries pool the sampled cropping treatments; the weather analysis does not estimate treatment effects. PAIC is neither standardised nor detrended before screening or modelling.')
p('The model dataset contains 26 measurement years from 1994 to 2023, giving 104 hybrid summaries. All candidate models use these same years, with complete summer temperature and rainfall. Weather is assigned from the growth year preceding measurement: a measurement in year t receives weather from year t − 1. There are 26 shared weather occasions, not 104 independent weather replicates.')
p('PAIC is divided by the years between successive measurements. This accounts for interval duration, but we do not know how growth was distributed among missing years. A multi-year PAIC value is still paired with the final growth year’s weather, following the agreed convention. For example, the 2001 PAIC averages growth years 1998–2000 and is paired with weather in 2000. This limits attribution of individual fluctuations to a particular year’s weather.')
h('Weather variables and biological hypotheses')
table(['Candidate group','Variables examined','Biological question'],[
    ['Water supply','Spring and summer rainfall totals (mm)','Does water availability coincide with greater growth?'],
    ['Warmth','Spring and summer mean air temperature (°C)','Does seasonal warmth coincide with greater growth?'],
    ['Light','Spring and summer sunshine totals (hours)','Does light availability coincide with greater growth?'],
    ['Dry conditions','Longest April–September dry spell (days)','Does sustained low rainfall coincide with reduced growth?'],
    ['Frost','Spring air and ground frost days','Does frost coincide with reduced growth?'],
], [1.05,3.0,2.55])
p('Spring means March–May and summer means June–August. A dry day has rainfall <1 mm; the dry spell is the longest consecutive run. Air and ground frost are days with minimum air and grass temperature below 0 °C, respectively. A candidate seasonal value is used only when all expected daily observations for that variable are available. No missing daily weather is imputed.')

page('Calculated results from the weather screen')
fig('raw_screen','Figure 1. Pearson correlations between each candidate and original PAIC, by hybrid, using 24 common complete-weather measurement years. The colour scale is centred on zero. These are descriptive associations and include shared temporal patterns.')
p('We inspected scatterplots for all nine variables within each hybrid and calculated Pearson correlations for linear association and Spearman correlations for rank association. The common 24-year comparison excludes measurement years 2001 and 2019 because of incomplete spring temperature or ground-frost records. Predictor-specific complete records (25 or 26 years) and omission of individual years were also checked; the final summer models use all 26 complete summer years.')
p('Summer rainfall has the clearest consistent positive raw association: Pearson r = 0.40, 0.56, 0.57 and 0.63 for Beaupre, Gibecq, Robusta and Trichobel. The corresponding Spearman values are 0.36, 0.48, 0.50 and 0.63. Every leave-one-year-out Pearson value remains positive. Summer temperature has weaker raw correlations: 0.03, 0.06, 0.20 and 0.27; Trichobel also has positive spring temperature and sunshine associations.')
h('How the model variables were chosen')
p('We retained summer rainfall as the clearest common water-supply signal and summer temperature as a biologically motivated warmth hypothesis, including the question of Trichobel’s response. The candidate family contains no weather, rainfall alone, temperature alone, and both together. This tests the added contribution of temperature rather than assuming it is needed. Spring, sunshine, frost and dry-spell variables remain contextual screening results; they were not added through stepwise searches.')
p('This is hypothesis-guided exploratory selection, informed by earlier examination of the dataset. It was not independently prespecified. Correlation strength alone did not determine inclusion, and the reported intervals do not account for this selection process.')

page('Statistical approach')
h('Allowing for stand development over time')
p('Original positive hybrid medians are fitted with a Gamma generalised linear model and a log link. This allows variability to increase with the fitted mean and keeps fitted PAIC positive. PAIC is not transformed, standardised or detrended before fitting. Each hybrid has its own intercept, weather slopes, age trajectory and dispersion estimate.')
p('For hybrid h and measurement year t, the combined-weather model is:')
p('PAICₕₜ ~ Gamma with mean μₕₜ\nlog(μₕₜ) = αₕ + βT,ₕ Tₜ + βR,ₕ Rₜ + γ₁,ₕ Aₜ + γ₂,ₕ Aₜ²')
p('T = summer mean temperature − 16 °C; R = (summer rainfall − 200 mm)/100; A = (age in years − 15)/10. Centring improves numerical stability. Both age and age squared are retained together; the squared term allows a curved trajectory on the log-mean scale, including a possible growth peak and subsequent decline.')
p('The age terms allow for PAIC changes associated with stand development over time, which can coincide with changing weather. This reduces the risk that the weather terms absorb the entire long-term growth pattern. All trees were planted together, but hybrids may have different growth trajectories. Age is a proxy for these temporal processes: it does not independently identify ageing or fully remove confounding by calendar year, management or stand conditions.')
p('The same quadratic age specification is used in all four candidate models: age alone, age plus summer rainfall, age plus summer temperature, and age plus both. This gives four fits per hybrid and 16 fits overall, all on the same 26 years. The combined-weather model supplies the comparable associations presented in Figure 2; rainfall-only and temperature-only candidates describe the contribution of the other weather term without a stepwise search.')
h('Absolute weather contrasts')
p('Temperature is contrasted from 15.31 to 16.31 °C, holding rainfall at 172.44 mm. Rainfall is contrasted from 122.44 to 222.44 mm, holding temperature at 15.81 °C. These contrasts are centred on observed means. Predictions are averaged equally over the same 26 observed ages. Each effect is the higher-weather prediction minus the lower-weather prediction, in Mg C ha⁻¹ yr⁻¹. It describes an absolute association at comparable modelled ages, not the change from the first to the final study year.')
h('What the dots and whiskers mean')
p('Dots are estimated mean PAIC differences for the stated weather contrasts. Whiskers show uncertainty in those estimates. We generated 1,000 alternative datasets by resampling blocks of three successive observed visits. All four hybrids and their weather move together. We refit the four hybrid models and recalculate the contrasts for each resample. The middle 95% of these values gives a pointwise percentile interval.')
p('Hybrid differences are calculated directly within each resample, preserving their pairing. A wider whisker means the mean association is less precisely estimated; it does not measure year-to-year variability in physiological sensitivity. The models do not estimate separate weather slopes for individual years. Blocks wrap across the record ends and cover three visits, not necessarily three calendar years. Intervals are exploratory, not simultaneous, and omit variable-selection, allometric and weather measurement uncertainty.')

page('Calculated weather associations by hybrid')
fig('hybrid_effects','Figure 2. Absolute PAIC differences associated with summer rainfall and temperature, conditional on hybrid-specific quadratic age trajectories. Dots show estimated mean differences and whiskers show pointwise 95% paired block-bootstrap intervals. Rainfall and temperature are included together. Whiskers quantify estimation uncertainty, not annual variation in the response.')
eff=rows('absolute_effects')
assert len(eff)==8 and all(e['age_model']=='Quadratic age' for e in eff)
body=[]
for hybrid in ['Beaupre','Gibecq','Robusta','Trichobel']:
    cells=[hybrid]
    for var in ['Rainfall','Temperature']:
        v=next(e for e in eff if e['hybrid']==hybrid and e['predictor']==var)
        cells.append(f"{float(v['difference']):.2f} ({float(v['lower']):.2f}, {float(v['upper']):.2f})")
    body.append(cells)
table(['Hybrid','100 mm more rainfall','1 °C higher temperature'],body,[1.25,2.65,2.7])
p('Table 1. Absolute mean PAIC differences (Mg C ha⁻¹ yr⁻¹) and pointwise 95% intervals from the combined-weather models. Contrasts are centred on mean weather and averaged over common observed ages. An interval extending below zero allows a negative association; fitted PAIC itself remains positive.','Caption')
p('Rainfall estimates are positive in all hybrids, and their intervals are above zero. Trichobel has the largest point estimate, 0.77 Mg C ha⁻¹ yr⁻¹ per 100 mm, followed by Beaupre (0.69), Gibecq (0.66) and Robusta (0.56). Direct paired comparisons do not clearly distinguish the hybrid effects: all their pointwise 95% intervals include zero.')
p('Temperature estimates are also positive. Trichobel has the largest point estimate, 0.58 Mg C ha⁻¹ yr⁻¹ per 1 °C, but its interval is broad and includes zero (−0.07 to 1.21). The other estimates are 0.44 for Beaupre, 0.39 for Gibecq and 0.53 for Robusta. None of the paired temperature differences between hybrids has an interval wholly above or below zero. The results therefore do not establish that Trichobel is either more or less temperature-sensitive than the others.')

page('Discussion and limits of interpretation')
h('Possible biological explanation')
p('The most consistent association is greater PAIC with wetter summers across all four hybrids, compatible with a contribution from summer water availability. Warmer summers are also associated with greater fitted PAIC, conditional on rainfall and the modelled age trajectory. These patterns provide a possible environmental contribution to growth during the study, rather than an independent attribution to climate change.')
p('Trichobel has the largest estimated absolute association with both rainfall and temperature. However, its temperature association is less precisely estimated, and direct comparisons do not establish that its weather responses differ from those of the other hybrids. Appropriate wording is: “Trichobel showed the largest estimated absolute PAIC response to increased summer rainfall, although differences between hybrids were uncertain.” Avoid describing the whiskers as evidence that its response varied more between years.')
p('A larger absolute effect can partly reflect Trichobel’s higher baseline PAIC: with a log link, the same proportional association produces more additional carbon at a higher baseline. Absolute differences answer how much carbon increment differs; they do not establish greater physiological sensitivity or climate resistance. The proposed resilience mechanism remains a discussion hypothesis supported by biological expertise, not demonstrated by these models.')
p('A positive weather association can coexist with declining PAIC through time. Observed growth reflects the actual combination of weather, age, stand development and management. Quadratic age allows a curved background trajectory, but age and calendar year advance together in this single cohort. The age terms may absorb weather-related trends as well as unrelated development, so they do not fully separate those processes.')
h('Uncertainty and diagnostic caveats')
p('Weather refers to the growth year preceding measurement. Annualisation accounts for the duration of multi-year PAIC intervals but cannot identify growth in missing years; using the final year’s weather leaves that timing uncertain. The 26 shared weather occasions are the basis for environmental inference; the four hybrid summaries do not multiply the independent weather sample size.')
p('Residual correlations between consecutive calendar years are approximately 0.35, 0.56, 0.29 and 0.00 for Beaupre, Gibecq, Robusta and Trichobel. Three-visit blocks preserve some local dependence, but their length is not established as sufficient. Interval endpoints remain exploratory. Stronger significance claims would require a defensible serial-error analysis or assessment of block-length sensitivity.')
audit=rows('bootstrap_tail_audit')
counts=[int(e['above_10_absolute']) for e in audit if int(e['above_10_absolute'])>0]
tail_range=f'{min(counts)}–{max(counts)}' if counts else '0'
p(f'All 4,000 bootstrap model fits converged. Nevertheless, some resamples produced very large absolute effects when evaluated across the full reference age range. For affected hybrid–predictor combinations, {tail_range} of 1,000 resamples exceeded an absolute difference of 10 Mg C ha⁻¹ yr⁻¹. These values were retained. The central percentile intervals are finite, but the tails indicate unstable extrapolation in some resamples and reinforce the exploratory interpretation. Individual-year omission results and tail checks are saved.')
h('Presentation in the paper')
p('Use Figure 2 for the weather associations and explain the hypothesis-guided variable selection in Methods. Figure 1 can document that selection in supplementary material. Figure S1 shows the fitted quadratic age trajectories at fixed weather, making the background age patterns explicit without treating them as isolated causal effects.')

page('Supplementary age trajectories and model fit')
fig('quadratic_age_shapes','Figure S1. Fitted mean PAIC across the observed age range, with summer temperature fixed at 15.81 °C and rainfall at 172.44 mm. Each hybrid has its own quadratic age terms. These curves illustrate model shapes; uncertainty is omitted for clarity. They are not observed growth trajectories or estimates of an isolated causal age effect.',width=6.2)
p('The fitted age trajectory declines in Beaupre, bends modestly in Gibecq, and allows a growth peak followed by decline in Robusta and Trichobel. These are conditional descriptions of this cohort at fixed weather, not evidence that hybrids would follow the same trajectory under other management or stand conditions.')
comparison=rows('candidate_models')
body=[]
for hybrid in ['Beaupre','Gibecq','Robusta','Trichobel']:
    v=next(x for x in comparison if x['hybrid']==hybrid and x['weather_model']=='Both')
    body.append([hybrid,f"{float(v['deviance_pseudo_R2']):.3f}",f"{100*float(v['deviance_reduction_vs_same_age']):.1f}%"])
table(['Hybrid','Combined model\npseudo R²','Weather deviance reduction\nagainst quadratic age alone'],body,[1.1,2.2,3.3])
p('Table S1. Pseudo-R² = 1 − combined-model deviance / intercept-only deviance and therefore describes age and weather together. The last column compares the combined model with quadratic age alone. These are improvements in fitted Gamma deviance, not percentages of raw PAIC variance explained or a causal share attributable to weather. AICc comparisons of all four weather candidates are exported separately and remain descriptive because residual serial dependence persists.','Caption')
p('Reproducibility: analysis/14_raw_paic_weather_final.R, its figures helper and 14_raw_paic_weather_checks.R; results/models/final_raw_paic_weather.rds; current final_weather_ tables and figures. Candidate weather values are reused from weather_departure_screen_data.csv; no departures from that file enter the models. Figures have 600 dpi PNG and vector PDF exports. Earlier plots are archived outside results/figures; earlier Word versions are preserved separately.','Caption')

d.core_properties.title='Weather associations with observed PAIC'
d.core_properties.subject='Quadratic age methods weather associations and supplementary trajectories'
d.save(OUT)
print(f'Saved {OUT}\nPreserved {backup}')
