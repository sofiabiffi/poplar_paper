from pathlib import Path
from datetime import datetime
import csv,shutil
from docx import Document
from docx.shared import Inches
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
d=Document(path);title='Hybrid specific summer weather associations'
if any(p.text==title for p in d.paragraphs): raise RuntimeError('Update already present.')
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'{path.stem}_before_hybrid_weather_{datetime.now():%Y%m%d_%H%M%S}.docx')
def read(n):
    with (root/f'results/tables/hybrid_weather_{n}.csv').open(newline='') as f:return list(csv.DictReader(f))
slopes=read('slopes');contrasts=read('paired_slope_differences');diagnostics=read('residual_summary')
d.add_page_break();d.add_heading(title,1)
d.add_paragraph('Update 20 October 2026. Summer rainfall and temperature have positive fitted associations with growth departures in all four hybrids. Trichobel does not have a clearly larger temperature association after removing broad growth trajectories. These exploratory results do not establish equal responses or causal weather effects.')
d.add_heading('Methods',2)
d.add_paragraph('Use the same 26 common measurement years and tree-level hybrid medians as the shared-departure models. Remove a separate quadratic age trajectory from log PAIC for each hybrid and standardise residuals by their within-hybrid standard deviation. Fit hybrid-specific intercepts and summer rainfall and temperature slopes, equivalently a fully interacting joint linear model. Directly compare paired slopes rather than comparing significance within hybrids. Standardised slopes describe sensitivity relative to each hybrid variability, not absolute differences in carbon increment.')
d.add_paragraph('For uncertainty, take 1000 paired circular moving-block bootstrap samples, each built from blocks of three ordered measurement visits. All four hybrids and their weather stay together. Refit age trends, residual scales and weather slopes in every sample. Blocks follow observed visits and can span missing-year gaps; the block length is provisional. Plot pointwise percentile intervals for slopes and calculate simultaneous approximate intervals over all 12 pairwise slope contrasts using the maximum centred, standardised bootstrap deviation. This is exploratory after full-data weather screening, not confirmatory testing.')
d.add_heading('Calculated results',2)
for h in ['Beaupre','Gibecq','Robusta','Trichobel']:
    r=next(x for x in slopes if x['hybrid']==h and x['term']=='rainfall');t=next(x for x in slopes if x['hybrid']==h and x['term']=='temperature')
    d.add_paragraph(f"{h}: rainfall slope {float(r['estimate']):.2f} standardised units per 100 mm; temperature slope {float(t['estimate']):.2f} per degree C; in-sample R squared {float(r['R2']):.2f}.")
excluded=[x for x in contrasts if x['interval_excludes_zero']=='TRUE']
d.add_paragraph(f"Of the 12 paired rainfall and temperature slope differences, {len(excluded)} have simultaneous bootstrap intervals excluding zero. These are direct slope comparisons, not comparisons of separate p values. Linear detrending also gives positive slopes in every hybrid and does not place Trichobel highest for temperature. Unstandardised log-departure slopes likewise do not show a larger Trichobel temperature coefficient.")
d.add_heading('Discussion hypotheses',2)
d.add_paragraph('The later relative advantage of Trichobel is visible in observed PAIC, but it does not translate into a clearly steeper summer temperature slope in the detrended analysis. The broad hybrid trajectories can absorb a gradual divergence that coincides with warming, so this result concerns fluctuations around those trajectories. It cannot determine whether warming contributed to the longer-term divergence. Age, time and cultivation history remain entangled in this planting cohort.')
d.add_heading('Open decisions and residual diagnostics',2)
d.add_paragraph('Weather models leave substantial temporal structure. Correlations between residuals in consecutive observed calendar years are about 0.59 for Beaupre, 0.68 for Gibecq, 0.59 for Robusta and 0.30 for Trichobel. Block-bootstrap uncertainty addresses some short dependence but does not model its cause. Assess block-length and age-trend sensitivity and unresolved shared conditions before paper inference. Omit-one-year slopes are exported as influence checks. The weather variable screening, multi-year interval, coverage, Robusta legacy discrepancy and PCA-retention caveats remain open. No forecast claim is intended.')
d.add_page_break();d.add_heading('Observed and fitted departures by hybrid',1)
d.add_picture(str(root/'results/figures/hybrid_weather_time_series_preview.png'),width=Inches(6.65))
d.add_paragraph('Observed departures and full-data weather fits show which peaks and troughs are captured. These are in-sample fits after removing age trajectories. The separate slope figure reports block-bootstrap intervals.',style='Caption')
d.add_heading('Reproducible sources',2)
d.add_paragraph('analysis/12_hybrid_weather_departures.R; results/tables/hybrid_weather_slopes.csv; hybrid_weather_paired_slope_differences.csv; hybrid_weather_sensitivity_slopes.csv; hybrid_weather_omit_year_slopes.csv; hybrid_weather_residual_summary.csv; results/models/hybrid_weather_departures.rds. Figures: hybrid_weather_time_series and hybrid_weather_slopes, PNG and PDF.')
d.save(path);print(path)
