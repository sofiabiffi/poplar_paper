from pathlib import Path
from datetime import datetime
import shutil
from docx import Document
root=Path(__file__).resolve().parents[1]
p=root/'manuscript/analysis_notes/04_soil_carbon_stocks.docx'
(p.parent/'_history').mkdir(parents=True, exist_ok=True); shutil.copy2(p,p.parent/'_history'/('04_soil_carbon_stocks_before_new_dataset_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.docx'))
d=Document(p)
d.add_heading('Replacement study site dataset',level=2)
d.add_paragraph('The new workbook outputs/soil_provenance_20261022/study_site_soil_2024.xlsx and matching data/processed/study_site_soil_2024.csv provide the replacement analysis input. The soil_0_30 sheet contains 188 measured layers from 96 cores, using full-precision company fine BD and SOC values. Corrected membership is 48 Arable reference cores (95 layers), 24 Alley cores (48 layers) and 24 Tree row cores (45 layers). Four cores lack the lower retained layer; no missing layers are imputed.')
d.add_paragraph('Sampling locations retain location_code values 1_tree, 2_alley and 3_arable with readable labels. Separate former_treatment_code values F, C and A identify Continuous fallow, Continuous cropping and Alternate cropping for both tree-row and alley cores. Block and original tree-row numbers are included. Former treatment is not a statement of management in 2024. Arable reference has no experimental cropping treatment, block or tree-row assignment. Permanent company core and sample IDs are the new join keys; original workbook IDs and labels remain in legacy audit columns. The company_site_results sheet preserves all 265 source records for the two study fields, including deeper samples, and Guide documents definitions, source provenance and limitations.')
d.add_paragraph('The dataset corrects legacy 1_tree_25 to Arable reference and preserves source precision. Existing SOC analysis scripts and numerical outputs have not yet been switched or rerun. Subsequent recalculation must use the new input and stable company IDs. Build with scripts/build_site_soil_data.py followed by scripts/build_site_soil_workbook.mjs; the original company results and scenario workbook remain unchanged.')
d.save(p)
