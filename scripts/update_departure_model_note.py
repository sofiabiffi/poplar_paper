from pathlib import Path
from datetime import datetime
import csv,shutil
from docx import Document
from docx.shared import Inches
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
d=Document(path);title='Weather models of shared PAIC departures'
if any(p.text==title for p in d.paragraphs): raise RuntimeError('Already appended; edit deliberately.')
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'{path.stem}_before_departure_models_{datetime.now():%Y%m%d_%H%M%S}.docx')
d.add_page_break();d.add_heading(title,1)
d.add_paragraph('Update 20 October 2026. Combined June to August rainfall and temperature improve prediction of shared standardised PAIC departures relative to no weather in this exploratory comparison. This is a different response and weather window from the earlier raw PAIC mixed model, and it is not an independent confirmation after screening.')
d.add_heading('Methods',2)
d.add_paragraph('Use 26 common measurement years with positive tree-level median PAIC for all four hybrids and complete summer rainfall and temperature coverage. In each training fold, fit a separate quadratic age trajectory to log PAIC for each hybrid. Divide training residuals by their within-hybrid training residual standard deviation. Calculate held-out departures from the training-fitted age trajectory and the same training standard deviation. Average the four departures with equal hybrid weights. No held-out PAIC enters the fitted trajectories or scales.')
d.add_paragraph('Fit linear models of shared departures using no weather, rainfall only, temperature only, or both. Age is not added to the second-stage regression. Rainfall is in 100 mm units and temperature in degrees C. Compare leave-one-year-out and five contiguous blocks of observed measurement years, keeping all hybrids from each year together. Repeat using linear rather than quadratic age detrending as a sensitivity. No automated predictor selection is performed inside this comparison; prior full-data screening remains a source of selection optimism. All multi-year increments retain the confirmed t minus 1 weather assignment.')
d.add_heading('Calculated results',2)
with (root/'results/tables/weather_departure_model_validation.csv').open(newline='') as f: rows=list(csv.DictReader(f))
for scheme in ['Leave one year out','Five contiguous blocks']:
    selected={r['model']:r for r in rows if r['degree']=='2' and r['scheme']==scheme}
    d.add_paragraph(f"{scheme}, quadratic detrending: no-weather RMSE {float(selected['Baseline']['RMSE']):.3f}; rainfall only {float(selected['Rainfall']['RMSE']):.3f}; temperature only {float(selected['Temperature']['RMSE']):.3f}; both {float(selected['Both']['RMSE']):.3f}. The combined model reduces RMSE by {float(selected['Both']['RMSE_improvement_percent']):.1f} percent relative to no weather within this scheme. Units are fold-standardised shared departures.")
d.add_paragraph('With linear detrending, the combined model reduces leave-one-year-out RMSE by 15.8 percent and block RMSE by 22.5 percent. Full-data combined-model slopes are 0.675 shared-departure units per 100 mm summer rainfall and 0.582 units per degree C. Its in-sample R squared is 0.305; this is not the proportion of raw PAIC causally explained by weather. Omitting 1994 from scoring only reduces the leave-one-year-out advantage: no-weather RMSE 0.957 versus combined-model 0.899, about 6 percent improvement; this is an influence check, not a refit excluding that year.')
d.add_heading('Discussion hypotheses',2)
d.add_paragraph('Warmer and wetter summer conditions are associated with more positive shared departures in the fitted range. The joint model performs better than either single predictor, suggesting potentially complementary information. These associations do not establish causation or imply that warming benefits growth outside observed conditions. The earlier failure of the half-year raw PAIC model and this result are not directly comparable: exposure window, response scale, data source and model structure differ.')
d.add_heading('Open decisions and limitations',2)
d.add_paragraph('Block errors are much larger than leave-one-year-out errors. The earliest block includes a held-out shared departure of about minus 9 training standard deviations, revealing poor extrapolation of the quadratic age trajectory. Do not interpret relative block improvement as reliable long-range weather forecasting. Full-data confidence bands condition on detrending and omit trend and predictor-selection uncertainty. The response is a shared standardised log-scale departure, not an individual hybrid PAIC forecast. Before inference, assess a less fragile growth trajectory, nested screening or independent data, residual temporal dependence, weather coverage and the legacy Robusta discrepancy. Keep sampling-gap and PCA-retention caveats open.')
d.add_page_break();d.add_heading('Summer weather models of shared departures',1)
d.add_picture(str(root/'results/figures/weather_departure_models_preview.png'),width=Inches(6.65))
d.add_paragraph('Curves show the combined model at the mean of the other predictor. Shading gives approximate 95 percent mean confidence intervals conditional on full-data detrending. Validation refits age trends and standardisation in each training fold; it does not undo prior full-data screening.',style='Caption')
d.add_heading('Reproducible sources',2)
d.add_paragraph('analysis/11_weather_departure_models.R; results/tables/weather_departure_model_validation.csv; weather_departure_model_held_out_predictions.csv; weather_departure_model_fold_audit.csv; weather_departure_model_coefficients.csv; results/models/weather_departure_models.rds.')
d.save(path);print(path)
