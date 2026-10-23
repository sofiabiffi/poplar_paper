from pathlib import Path
from datetime import datetime
import csv,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
root=Path(__file__).resolve().parents[1]
path=root/'manuscript/analysis_notes/02_weather_and_PAIC.docx'
d=Document(path)
title='Weather associations with annual carbon increment'
if any(p.text==title for p in d.paragraphs):raise RuntimeError('Summary already exists; preserve edits and update deliberately.')
history=path.parent/'_history';history.mkdir(exist_ok=True)
shutil.copy2(path,history/f'{path.stem}_before_paper_summary_{datetime.now():%Y%m%d_%H%M%S}.docx')
with (root/'results/tables/absolute_weather_adjusted_differences.csv').open(newline='') as f:rows=list(csv.DictReader(f))
d.add_page_break()
start=len(d.paragraphs)
p=d.add_paragraph(title,'Heading 1');p.paragraph_format.space_before=Pt(0)
d.add_paragraph('Paper analysis summary | 21 October 2026. Focus on absolute and percentage PAIC associations; retain standardised analyses as exploratory background only.')
def section(name,text):
    p=d.add_paragraph(name,'Heading 2');p.paragraph_format.space_before=Pt(6);p.paragraph_format.space_after=Pt(3)
    d.add_paragraph(text)
section('Weather variable selection',
    'We formulated hypotheses about seasonal warmth, water supply, sunshine, frost and dry spells. Exploratory correlations with growth departures highlighted summer rainfall and temperature. Both were retained in the final candidate model without further screening. Selection was informed by the same data used for fitting, so results are exploratory associations, not independent confirmation of the hypotheses.')
section('Response and previous year weather',
    'The analysis uses positive tree-level median PAIC for each of four hybrids in 26 measurement years (104 summaries). PAIC is carbon increment divided by years between measurements, in Mg C ha-1 yr-1. Each measurement in year t is paired with June-August mean air temperature and total rainfall in year t minus 1. For example, 2002 PAIC uses summer 2001 weather. For multi-year intervals, annualisation accounts for duration but cannot reveal growth timing in the missing years; weather still describes only the final growth year.')
section('Statistical approach and uncertainty',
    'Gamma models with a log link were fitted to original PAIC, allowing each hybrid its own baseline, curved age trajectory (age and age squared), and temperature and rainfall associations. Absolute differences average predictions over the same observed ages, holding the other weather variable at its mean. We compare 15.31 with 16.31 degrees C, and 122.4 with 222.4 mm rainfall. Percentage differences express the same comparisons relative to the lower-weather prediction.')
d.add_paragraph('To assess uncertainty, we repeatedly resampled blocks of three neighbouring visits, keeping all hybrids together, and refitted the complete model 1,000 times. Some visits were repeated and others omitted. The middle 95% of estimates gives the reported intervals. This preserves some shared and temporal variation, but blocks span sampling gaps and wrap from the record end to its start. Rare samples produced extreme age-related estimates; intervals remain provisional.')
section('Model results','Adjusted differences below are in Mg C ha-1 yr-1, with percentage differences in brackets. They describe fitted associations within observed weather conditions, not causal warming or rainfall scenarios.')
t=d.add_table(rows=1,cols=3);t.autofit=False
widths=[1.1,2.7,2.7]
for i,(text,w) in enumerate(zip(['Hybrid','Temperature difference of 1 degree C','Rainfall difference of 100 mm'],widths)):
    t.columns[i].width=Inches(w);t.rows[0].cells[i].text=text
for h in ['Beaupre','Gibecq','Robusta','Trichobel']:
    temp=next(x for x in rows if x['hybrid']==h and x['term']=='temperature');rain=next(x for x in rows if x['hybrid']==h and x['term']=='rainfall')
    cells=t.add_row().cells
    for c,text in zip(cells,[h,f"+{float(temp['absolute_difference']):.2f} (+{float(temp['percent_difference']):.1f}%)",f"+{float(rain['absolute_difference']):.2f} (+{float(rain['percent_difference']):.1f}%)"]):c.text=text
for ri,row in enumerate(t.rows):
    for ci,c in enumerate(row.cells):
        c.width=Inches(widths[ci]);pr=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        pr.append(borders)
        if ri==0:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF0');pr.append(shade)
        for p in c.paragraphs:
            p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
            for r in p.runs:r.font.size=Pt(9);r.bold=ri==0
d.add_paragraph('Trichobel has the largest absolute point estimates but the smallest percentage estimates. Direct paired comparisons do not clearly distinguish hybrids: all 12 difference intervals include zero. Trichobel temperature association is particularly uncertain (absolute 95% interval -0.09 to +1.22).')
section('Discussion and limits',
    'The positive associations offer a possible contribution to the observed growth patterns and are compatible with the biological expectation that Trichobel copes well with changing conditions. Its larger absolute estimate partly reflects higher growth, rather than a larger proportional response. All trees belong to one planting cohort: age, calendar time, management history and changing weather cannot be cleanly separated. Residual temporal patterns and bootstrap instability also remain. Present this as a supporting, exploratory explanation of Trichobel performance, without attributing its study-period advantage to climate change or claiming superior heat tolerance.')
p=d.add_paragraph('Sources: analysis/13_absolute_weather_paic.R; results/tables/absolute_weather_adjusted_differences.csv and absolute_weather_paired_differences.csv. Full intervals and figures appear in the preceding note and saved outputs.')
p.paragraph_format.space_after=Pt(0)
for p in d.paragraphs[start:]:
    if p.style.name not in ['Heading 1','Heading 2']:
        p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.0
        for r in p.runs:r.font.size=Pt(9.5)
    for r in p.runs:r.font.color.rgb=RGBColor(0,0,0)
d.save(path);print(path)
